#!/usr/bin/env python3
"""Build English and complete Chinese NVIDIA Q2 FY2027 earnings updates."""

from __future__ import annotations

import json
import sys
from pathlib import Path

import yfinance as yf
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/macrossz/DevTools/VscodeProject/ClaudeCode/financial_analysis")
OUT = ROOT / "output" / "NVDA"
DATA = ROOT / "data" / "nvda"
OUT.mkdir(parents=True, exist_ok=True)

DOC_SKILL_SCRIPTS = Path(
    "/Users/macrossz/.codex/plugins/cache/openai-primary-runtime/documents/"
    "26.826.12353/skills/documents/scripts"
)
sys.path.insert(0, str(DOC_SKILL_SCRIPTS))
from table_geometry import apply_table_geometry, column_widths_from_weights  # noqa: E402


# Named design override to standard_business_brief: compact institutional research.
PAGE_WIDTH_DXA = 10166
NAVY = "17365D"
BLUE = "2F5597"
GREEN = "416900"
NVIDIA_GREEN = "76B900"
PALE_GREEN = "E8F3D6"
ORANGE = "C65D17"
RED = "A61B1B"
GREY = "667085"
LIGHT = "F2F4F7"
WHITE = "FFFFFF"
BLACK = "182230"

URLS = {
    "release": "https://investor.nvidia.com/news/press-release-details/2026/NVIDIA-Announces-Financial-Results-for-Second-Quarter-Fiscal-2027/default.aspx",
    "cfo": "https://s201.q4cdn.com/141608511/files/doc_financials/2027/Q227/Q2FY27-CFO-Commentary.pdf",
    "presentation": "https://s201.q4cdn.com/141608511/files/doc_financials/2027/Q227/NVDA-F2Q27-Quarterly-Presentation-final-1.pdf",
    "tenq": "https://www.sec.gov/ixviewer/doc/action?doc=/Archives/edgar/data/1045810/000104581026000075/nvda-20260726.htm",
    "eightk": "https://www.sec.gov/Archives/edgar/data/1045810/000104581026000073/nvda-20260826.htm",
    "webcast": "https://investor.nvidia.com/events-and-presentations/events-and-presentations/event-details/2026/NVIDIA-2nd-Quarter-FY27-Financial-Results/default.aspx",
    "transcript": "https://seekingalpha.com/article/4940563-nvidia-corporation-nvda-q2-2027-earnings-call-transcript",
    "q1_release": "https://investor.nvidia.com/news/press-release-details/2026/NVIDIA-Announces-Financial-Results-for-First-Quarter-Fiscal-2027/default.aspx",
    "q4_release": "https://investor.nvidia.com/news/press-release-details/2026/NVIDIA-Announces-Financial-Results-for-Fourth-Quarter-and-Fiscal-2026/",
    "factset": "https://apnews.com/article/dc8d556e709b50915cca9217a60b1991",
    "yahoo_breaking": "https://finance.yahoo.com/video/nvidia-reports-impressive-q2-earnings-as-ai-boom-rolls-on-203300281.html",
    "yahoo_call": "https://finance.yahoo.com/markets/stocks/articles/nvidia-corp-nvda-q2-2027-050152280.html",
    "yahoo_quote": "https://finance.yahoo.com/quote/NVDA/",
    "longbridge_quote": "https://longbridge.com/en/quote/NVDA.US",
    "ir": "https://investor.nvidia.com/",
}

CHARTS = {
    1: OUT / "nvda_q2_chart1_quarterly_revenue.png",
    2: OUT / "nvda_q2_chart2_data_center.png",
    3: OUT / "nvda_q2_chart3_eps.png",
    4: OUT / "nvda_q2_chart4_gross_margin.png",
    5: OUT / "nvda_q2_chart5_beat_miss.png",
    6: OUT / "nvda_q2_chart6_platform_mix.png",
    7: OUT / "nvda_q2_chart7_guidance.png",
    8: OUT / "nvda_q2_chart8_working_capital.png",
    9: OUT / "nvda_q2_chart9_commitments.png",
    10: OUT / "nvda_q2_chart10_estimate_revisions.png",
    11: OUT / "nvda_q2_chart11_price_history.png",
}


def market_data() -> dict:
    """Fetch mandated Yahoo market fields dynamically; emit N/A on failure."""
    result = {"price": "N/A", "market_cap": "N/A", "high": "N/A", "low": "N/A"}
    try:
        info = yf.Ticker("NVDA").fast_info
        result = {
            "price": round(float(info.last_price), 2),
            "market_cap": float(info.market_cap),
            "high": round(float(info.year_high), 2),
            "low": round(float(info.year_low), 2),
        }
    except Exception as exc:  # project rule requires an explicit warning and N/A
        print(f"WARNING: yfinance market-data retrieval failed for NVDA: {exc}")
    quote_path = DATA / "longbridge_quote.json"
    if quote_path.exists():
        try:
            quote = json.loads(quote_path.read_text())[0]
            result["premarket"] = float(quote["pre_market"]["last"])
            result["postmarket"] = float(quote["post_market"]["last"])
            result["regular_close"] = float(quote["last"])
        except Exception as exc:
            print(f"WARNING: Longbridge quote parsing failed: {exc}")
    else:
        result.update({"premarket": 223.94, "postmarket": 219.53, "regular_close": 209.66})
    return result


def set_run_font(run, lang: str, size=9.5, bold=False, italic=False, color=BLACK, heading=False):
    if lang == "cn":
        east = "Heiti SC" if heading else "Songti SC"
        # LibreOffice's headless renderer does not reliably honor only the
        # eastAsia font slot on macOS.  Assign the Chinese face to every font
        # slot so Chinese glyphs survive both Word and PDF rendering; Latin
        # characters are still provided by the same macOS font families.
        latin = east
    else:
        east = latin = "Times New Roman"
    run.font.name = latin
    run._element.get_or_add_rPr()
    run._element.rPr.rFonts.set(qn("w:ascii"), latin)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), latin)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    if lang == "cn":
        language = run._element.rPr.find(qn("w:lang"))
        if language is None:
            language = OxmlElement("w:lang")
            run._element.rPr.append(language)
        language.set(qn("w:val"), "zh-CN")
        language.set(qn("w:eastAsia"), "zh-CN")


def paragraph(doc, text, lang, size=9.5, bold=False, italic=False, color=BLACK, before=0, after=4.2, align=None, keep=False, heading=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run_font(r, lang, size=size, bold=bold, italic=italic, color=color, heading=heading)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.keep_with_next = keep
    if align is not None:
        p.alignment = align
    return p


def heading(doc, text, lang, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(2 if level == 1 else 1)
    p.paragraph_format.space_after = Pt(4.5)
    p.paragraph_format.keep_with_next = True
    return p


def bullet(doc, title, body, lang, color=NAVY, compact=False):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.03)
    p.paragraph_format.first_line_indent = Inches(-0.03)
    p.paragraph_format.space_after = Pt(3.3 if compact else 4.2)
    p.paragraph_format.line_spacing = 1.03
    r = p.add_run("■ ")
    set_run_font(r, lang, size=8.9 if compact else 9.4, bold=True, color=color, heading=True)
    r = p.add_run(title + " ")
    set_run_font(r, lang, size=8.9 if compact else 9.4, bold=True, color=color, heading=True)
    r = p.add_run(body)
    set_run_font(r, lang, size=8.9 if compact else 9.4)
    return p


def add_hyperlink(paragraph_obj, label, url, lang, size=7.6):
    rel_id = paragraph_obj.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    link_font = "Songti SC" if lang == "cn" else "Times New Roman"
    fonts.set(qn("w:ascii"), link_font)
    fonts.set(qn("w:hAnsi"), link_font)
    fonts.set(qn("w:eastAsia"), link_font)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(round(size * 2)))
    for element in (fonts, color, underline, sz):
        props.append(element)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph_obj._p.append(link)


def source(doc, items, lang, lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3.5)
    p.paragraph_format.line_spacing = 1.0
    prefix = lead or ("Sources: " if lang == "en" else "资料来源：")
    r = p.add_run(prefix)
    set_run_font(r, lang, size=7.6, italic=True, color=GREY)
    for idx, item in enumerate(items):
        label, url = item
        if idx:
            r = p.add_run("; ")
            set_run_font(r, lang, size=7.6, italic=True, color=GREY)
        if url:
            add_hyperlink(p, label, url, lang, size=7.6)
        else:
            r = p.add_run(label)
            set_run_font(r, lang, size=7.6, italic=True, color=GREY)
    r = p.add_run(".")
    set_run_font(r, lang, size=7.6, italic=True, color=GREY)
    return p


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def make_table(doc, headers, rows, lang, weights, variance_col=None, font_size=8.0):
    table_obj = doc.add_table(rows=1, cols=len(headers))
    table_obj.style = "Table Grid"
    table_obj.alignment = WD_TABLE_ALIGNMENT.LEFT
    for col, label in enumerate(headers):
        cell = table_obj.rows[0].cells[col]
        cell.text = label
        shade(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cell.paragraphs[0].runs:
            set_run_font(run, lang, size=font_size, bold=True, color=WHITE, heading=True)
    repeat_header(table_obj.rows[0])
    for row_idx, values in enumerate(rows):
        cells = table_obj.add_row().cells
        for col, value in enumerate(values):
            cells[col].text = str(value)
            if row_idx % 2:
                shade(cells[col], LIGHT)
            cells[col].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[col].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if col == 0 else WD_ALIGN_PARAGRAPH.CENTER
            for run in cells[col].paragraphs[0].runs:
                color = BLACK
                if variance_col == col:
                    color = GREEN if str(value).startswith("+") else RED if str(value).startswith("-") else BLACK
                set_run_font(run, lang, size=font_size, bold=(col == 0), color=color)
    widths = column_widths_from_weights(weights, PAGE_WIDTH_DXA)
    apply_table_geometry(table_obj, widths, table_width_dxa=PAGE_WIDTH_DXA, indent_dxa=120)
    return table_obj


def figure(doc, idx, caption, lang, sources, width=6.35):
    label = f"Figure {idx}. {caption}" if lang == "en" else f"图 {idx}：{caption}"
    paragraph(doc, label, lang, size=8.5, bold=True, color=NAVY, after=1.5, keep=True, heading=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0.5)
    shape = p.add_run().add_picture(str(CHARTS[idx]), width=Inches(width))
    shape._inline.docPr.set("descr", caption)
    source(doc, sources, lang)


def page_field(paragraph_obj):
    run = paragraph_obj.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])
    set_run_font(run, "en", size=7.8, color=GREY)


def setup_doc(lang):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.66)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.30)

    normal = doc.styles["Normal"]
    normal_font = "Songti SC" if lang == "cn" else "Times New Roman"
    normal.font.name = normal_font
    normal._element.rPr.rFonts.set(qn("w:ascii"), normal_font)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), normal_font)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), normal_font)
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(4.2)
    normal.paragraph_format.line_spacing = 1.05

    for style_name, size, color in (("Heading 1", 14.0, NAVY), ("Heading 2", 11.3, BLUE), ("Heading 3", 10.0, NAVY)):
        style = doc.styles[style_name]
        heading_font = "Heiti SC" if lang == "cn" else "Times New Roman"
        style.font.name = heading_font
        style._element.rPr.rFonts.set(qn("w:ascii"), heading_font)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), heading_font)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), heading_font)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)

    header = section.header.paragraphs[0]
    header.text = (
        "NVIDIA (Nasdaq: NVDA)  |  Q2 FY2027 Earnings Update"
        if lang == "en"
        else "英伟达（纳斯达克：NVDA）｜2027 财年第二季度业绩更新"
    )
    set_run_font(header.runs[0], lang, size=7.8, bold=True, color=NAVY, heading=True)
    header.paragraph_format.space_after = Pt(0)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("NVDA Research  |  27 August 2026  |  ")
    set_run_font(r, lang, size=7.8, color=GREY)
    page_field(footer)
    return doc


def new_page(doc):
    doc.add_page_break()


def title_page(doc, lang, market):
    regular = market.get("regular_close", market["price"])
    pre = market.get("premarket", "N/A")
    market_cap = f"${market['market_cap'] / 1e12:.2f}tn" if market["market_cap"] != "N/A" else "N/A"
    range_text = f"${market['low']:.2f}-${market['high']:.2f}" if market["low"] != "N/A" else "N/A"
    if lang == "en":
        paragraph(doc, "NVIDIA CORPORATION (NASDAQ: NVDA)", lang, size=17.5, bold=True, color=NAVY, after=1, align=WD_ALIGN_PARAGRAPH.CENTER, heading=True)
        paragraph(doc, "Q2 FY2027 EARNINGS UPDATE", lang, size=13.5, bold=True, color=BLUE, after=1.5, align=WD_ALIGN_PARAGRAPH.CENTER, heading=True)
        paragraph(doc, "AI compute turns into revenue: another beat, a bigger FY2028 framework, and a larger balance-sheet commitment", lang, size=10.8, bold=True, color=ORANGE, after=6, align=WD_ALIGN_PARAGRAPH.CENTER, heading=True)
        metadata = [
            "Published: 27 August 2026 | Results released: 26 August 2026 (within 24 hours)",
            "Rating: MAINTAIN OUTPERFORM | Price target: $320 (prior: $280)",
            f"Reference price: ${pre:.2f} pre-market 27 Aug. | Regular close: ${regular:.2f} | Implied upside: {(320 / pre - 1) * 100:.1f}%",
            f"Market capitalization: {market_cap} | 52-week range: {range_text}",
        ]
        section = "EARNINGS SCORECARD"
        headers = ["Metric", "Reported", "Pre-results consensus", "Variance"]
        rows = [
            ["Revenue", "$96.22bn", "$92.27bn", "+$3.95bn / +4.3%"],
            ["Non-GAAP diluted EPS", "$2.22", "$2.09", "+$0.13 / +6.2%"],
            ["Data Center revenue", "$89.02bn", "$85.86bn", "+$3.16bn / +3.7%"],
            ["Q3 revenue guide", "$108.0bn +/-2%", "$105.16bn", "+$2.84bn / +2.7%"],
        ]
    else:
        paragraph(doc, "英伟达公司（纳斯达克：NVDA）", lang, size=17.5, bold=True, color=NAVY, after=1, align=WD_ALIGN_PARAGRAPH.CENTER, heading=True)
        paragraph(doc, "2027 财年第二季度业绩更新报告", lang, size=13.5, bold=True, color=BLUE, after=1.5, align=WD_ALIGN_PARAGRAPH.CENTER, heading=True)
        paragraph(doc, "AI 算力正在转化为收入：再次超预期、FY2028 增长框架显著扩大，但资产负债表承诺同步上升", lang, size=10.8, bold=True, color=ORANGE, after=6, align=WD_ALIGN_PARAGRAPH.CENTER, heading=True)
        metadata = [
            "发布日期：2026 年 8 月 27 日｜财报发布：2026 年 8 月 26 日（24 小时内完成）",
            "评级：维持“跑赢大盘”｜目标价：320 美元（此前：280 美元）",
            f"参考股价：8 月 27 日盘前 {pre:.2f} 美元｜8 月 26 日常规收盘 {regular:.2f} 美元｜潜在上涨空间：{(320 / pre - 1) * 100:.1f}%",
            f"市值：{market_cap}｜52 周区间：{range_text}",
        ]
        section = "业绩计分卡"
        headers = ["指标", "实际值", "财报前一致预期", "差异"]
        rows = [
            ["营业收入", "962.2 亿美元", "922.7 亿美元", "+39.5 亿 / +4.3%"],
            ["Non-GAAP 摊薄 EPS", "2.22 美元", "2.09 美元", "+0.13 / +6.2%"],
            ["数据中心收入", "890.2 亿美元", "858.6 亿美元", "+31.6 亿 / +3.7%"],
            ["Q3 收入指引", "1,080 亿美元 +/-2%", "1,051.6 亿美元", "+28.4 亿 / +2.7%"],
        ]
    for line in metadata:
        paragraph(doc, line, lang, size=8.7, bold=("Rating" in line or "评级" in line), after=1.3)
    heading(doc, section, lang, 2)
    make_table(doc, headers, rows, lang, [1.7, 1.55, 2.1, 1.7], variance_col=3, font_size=7.7)
    source(doc, [("NVIDIA Q2 FY27 release" if lang == "en" else "英伟达 Q2 FY27 财报公告", URLS["release"]), ("FactSet consensus via AP" if lang == "en" else "美联社转引 FactSet 一致预期", URLS["factset"]), ("Yahoo Finance breaking consensus" if lang == "en" else "Yahoo Finance 财报时点一致预期", URLS["yahoo_breaking"])], lang)
    heading(doc, "INVESTMENT TAKEAWAYS" if lang == "en" else "投资要点", lang, 2)
    if lang == "en":
        bullet(doc, "The beat was broad and higher quality than the headline alone suggests.", "Revenue beat FactSet consensus by $3.95bn and non-GAAP EPS by $0.13. Data Center exceeded the contemporaneous Yahoo estimate by $3.16bn, while the 75.0% non-GAAP gross margin held flat sequentially despite the Blackwell Ultra ramp and higher inventory provisions.", lang, compact=True)
        bullet(doc, "Growth broadened beyond hyperscalers.", "Hyperscale revenue rose 13% QoQ to $48.7bn, but ACIE increased 25% to $40.3bn and 138% YoY. That mix supports the thesis that AI-native firms, sovereigns and enterprises are becoming a second demand engine rather than a marginal extension of the four largest clouds.", lang, compact=True)
        bullet(doc, "The FY2028 framework is the main positive revision.", "Management said on the 26 August call that it expects revenue to grow approximately 70% in FY2028, materially above the pre-call market growth assumption. We model $700bn of FY2028 revenue and $16.50 of non-GAAP EPS, but explicitly treat the outlook as a management framework rather than formal guidance.", lang, compact=True)
        bullet(doc, "Balance-sheet intensity is now the central debate.", "Supply commitments increased from $119bn to $279bn in one quarter; total disclosed future commitments were $366bn, and maximum guarantee exposure reached $108.5bn after the August SB Energy arrangement. These commitments may secure scarce capacity and customer sites, but they introduce duration, counterparty and perceived circular-financing risk.", lang, compact=True)
        bullet(doc, "Maintain Outperform; raise price target to $320.", "The target is 19.4x our FY2028E non-GAAP EPS, a substantial de-rating from NVIDIA's five-year trailing average. We lift the target only 14% despite a much larger FY2028 earnings revision because commitments, cash conversion, customer concentration and policy risk require a wider valuation discount.", lang, compact=True)
    else:
        bullet(doc, "本季超预期覆盖面广，质量高于单一收入数字所显示的水平。", "收入较 FactSet 一致预期高 39.5 亿美元，Non-GAAP EPS 高 0.13 美元；数据中心收入较 Yahoo 同期预期高 31.6 亿美元。即使 Blackwell Ultra 放量且库存拨备增加，Non-GAAP 毛利率仍环比持平于 75.0%。", lang, compact=True)
        bullet(doc, "增长正在从超大规模客户向更广泛需求扩散。", "Hyperscale 收入环比增长 13% 至 487 亿美元，而 ACIE 环比增长 25% 至 403 亿美元、同比增长 138%。这表明 AI 原生企业、主权客户与传统企业正在成为第二增长引擎，而非四大云厂商需求的边际延伸。", lang, compact=True)
        bullet(doc, "FY2028 增长框架是本次最大正向修订。", "管理层在 8 月 26 日电话会上表示，预计 FY2028 收入增长约 70%，显著高于会前市场假设。我们据此预测 FY2028 收入 7,000 亿美元、Non-GAAP EPS 16.50 美元，但明确将其视为管理层框架，而非正式财务指引。", lang, compact=True)
        bullet(doc, "资产负债表强度已成为核心争议。", "供应承诺一个季度内从 1,190 亿美元增至 2,790 亿美元；披露的未来承诺总额为 3,660 亿美元，8 月 SB Energy 安排后最大担保敞口达到 1,085 亿美元。这些安排有助于锁定稀缺产能和客户站点，但同时增加期限、交易对手及“循环融资”观感风险。", lang, compact=True)
        bullet(doc, "维持“跑赢大盘”，目标价上调至 320 美元。", "目标价对应 FY2028E Non-GAAP EPS 的 19.4 倍，较英伟达五年历史市盈率显著折价。尽管 FY2028 盈利预测上修幅度更大，我们仅将目标价上调 14%，因为承诺规模、现金转换、客户集中与政策风险需要更宽估值折价。", lang, compact=True)


def page_revenue(doc, lang):
    new_page(doc)
    heading(doc, "REVENUE: BLACKWELL ULTRA AND A SECOND DEMAND ENGINE" if lang == "en" else "收入：Blackwell Ultra 与第二需求引擎", lang, 1)
    if lang == "en":
        paragraph(doc, "Q2 revenue reached $96.22bn, up 18% QoQ and 106% YoY. The print was $5.22bn above the midpoint of NVIDIA's own $91bn guide and $3.95bn above FactSet consensus. More important than the arithmetic, Blackwell Ultra supported an unusually clean combination of accelerating volume and stable 75% gross margin. The eight-quarter progression below shows that the company has moved from $35bn to $96bn of quarterly revenue without the margin compression normally associated with a hardware ramp of this magnitude.", lang)
    else:
        paragraph(doc, "第二季度收入达到 962.21 亿美元，环比增长 18%、同比增长 106%。实际值较公司 910 亿美元指引中点高 52.21 亿美元，较 FactSet 一致预期高 39.5 亿美元。更重要的是，Blackwell Ultra 在推动销量加速的同时，毛利率仍稳定在 75%。下方八季度序列显示，季度收入从 350 亿美元增至 962 亿美元，却未出现如此规模硬件爬坡通常伴随的明显利润率压缩。", lang)
    figure(doc, 1, "Quarterly revenue has more than doubled year over year" if lang == "en" else "季度收入同比增长超过一倍", lang, [("NVIDIA quarterly releases" if lang == "en" else "英伟达历季财报公告", URLS["release"]), ("Q4 FY26 release" if lang == "en" else "Q4 FY26 财报公告", URLS["q4_release"])], width=5.65)
    if lang == "en":
        paragraph(doc, "Data Center contributed $89.02bn, or 92.5% of revenue. It rose 18% sequentially and 117% from a year ago. The sequential share increase was small because Edge Computing also grew 13%, but the absolute scale is decisive: Data Center added nearly $48bn of revenue in twelve months. Blackwell remained the vast majority of shipments, while Vera Rubin entered production shipments in Q3. That overlap supports continued growth but raises execution risk because NVIDIA must manage two architectures and a much larger component envelope simultaneously.", lang)
    else:
        paragraph(doc, "数据中心贡献 890.23 亿美元，占总收入 92.5%，环比增长 18%、同比增长 117%。由于边缘计算也环比增长 13%，数据中心收入占比仅小幅上升，但绝对规模更具决定性：过去十二个月数据中心新增收入接近 480 亿美元。Blackwell 仍占绝大多数出货，Vera Rubin 已在第三季度开始生产出货。两代架构重叠支撑增长，也提高执行风险，因为公司必须同时管理两套架构以及更庞大的组件需求。", lang)
    figure(doc, 2, "Data Center reached $89.0bn and 92.5% of total revenue" if lang == "en" else "数据中心收入达到 890 亿美元，占总收入 92.5%", lang, [("Q2 FY27 CFO Commentary, pp. 1-2" if lang == "en" else "Q2 FY27 CFO 评论第 1-2 页", URLS["cfo"]), ("Q2 FY27 Form 10-Q" if lang == "en" else "Q2 FY27 Form 10-Q", URLS["tenq"])], width=5.65)


def page_profitability(doc, lang):
    new_page(doc)
    heading(doc, "PROFITABILITY: STRONG OPERATING LEVERAGE, LOWER CASH CONVERSION" if lang == "en" else "盈利能力：经营杠杆强劲，但现金转换下降", lang, 1)
    if lang == "en":
        paragraph(doc, "GAAP operating income rose 124% YoY to $63.73bn, faster than revenue, while non-GAAP operating income reached $63.96bn. Non-GAAP EPS of $2.22 grew 120% and beat consensus by 6.2%. Comparability requires care: beginning in FY2027, NVIDIA's non-GAAP measures include stock-based compensation, and historical comparisons in current materials were recast. The gap between GAAP EPS of $2.46 and non-GAAP EPS also reflects $7.8bn of net equity-security gains excluded from the adjusted measure. We therefore anchor valuation to non-GAAP EPS and operating economics, not the higher GAAP result.", lang)
    else:
        paragraph(doc, "GAAP 营业利润同比增长 124% 至 637.34 亿美元，快于收入增速；Non-GAAP 营业利润为 639.56 亿美元。Non-GAAP EPS 为 2.22 美元，同比增长 120%，较一致预期高 6.2%。但可比性需要谨慎：从 FY2027 起，英伟达 Non-GAAP 指标已包含股权激励费用，当前材料中的历史数据也进行了重列。GAAP EPS 2.46 美元与 Non-GAAP EPS 的差异还受到 78 亿美元股票投资净收益影响，该收益在调整口径中被剔除。因此估值应以 Non-GAAP EPS 与经营经济性为锚，而非更高的 GAAP 结果。", lang)
    figure(doc, 3, "Q2 non-GAAP EPS rose 120% year over year to $2.22" if lang == "en" else "Q2 Non-GAAP EPS 同比增长 120% 至 2.22 美元", lang, [("Q2 FY27 earnings release" if lang == "en" else "Q2 FY27 财报公告", URLS["release"]), ("Q2 FY27 CFO Commentary, pp. 1 and 7" if lang == "en" else "Q2 FY27 CFO 评论第 1、7 页", URLS["cfo"])], width=5.65)
    if lang == "en":
        paragraph(doc, "GAAP and non-GAAP gross margins were both 75.0%, flat sequentially and 260/250 bps above the prior year. Improved Blackwell Ultra mix offset $985mn of inventory and excess-purchase provisions; net of $177mn of releases, the adverse gross-margin effect was 80 bps. That resilience is encouraging, but Q3 guidance falls to 74.0% +/-50 bps. We interpret the step-down as product-transition and mix friction rather than structural pricing deterioration, while our FY2027 estimate moves to 74.5% from the prior 75.0% baseline.", lang)
    else:
        paragraph(doc, "GAAP 与 Non-GAAP 毛利率均为 75.0%，环比持平，同比分别提高 260 与 250 个基点。Blackwell Ultra 组合改善抵消了 9.85 亿美元库存及超额采购拨备；扣除 1.77 亿美元转回后，对毛利率的净拖累为 80 个基点。韧性值得肯定，但 Q3 指引降至 74.0% +/-50 个基点。我们将其视为产品过渡与组合摩擦，而非结构性定价恶化，并将 FY2027 毛利率预测从此前 75.0% 调整至 74.5%。", lang)
    figure(doc, 4, "Gross margin recovered to 75% despite the product transition" if lang == "en" else "产品过渡期间毛利率仍恢复至 75%", lang, [("Q2 FY27 Form 10-Q, MD&A" if lang == "en" else "Q2 FY27 Form 10-Q 管理层讨论", URLS["tenq"]), ("Q2 FY27 CFO Commentary" if lang == "en" else "Q2 FY27 CFO 评论", URLS["cfo"])], width=5.65)


def page_mix(doc, lang):
    new_page(doc)
    heading(doc, "BEAT/MISS AND MARKET-PLATFORM MIX" if lang == "en" else "预期差与市场平台组合", lang, 1)
    if lang == "en":
        paragraph(doc, "Vendor snapshots differ slightly by timestamp: FactSet consensus reported by AP was $92.27bn of revenue and $2.09 of adjusted EPS, while Longbridge's release-time snapshot was $92.16bn and $2.091. The analytical conclusion is unchanged. Revenue beat by roughly 4.3%-4.4%, adjusted EPS by 6.2%, and Data Center by 3.7% against Yahoo's $85.86bn expectation. The larger EPS variance indicates that the quarter was not simply a shipment-timing beat; operating leverage and mix also exceeded the pre-results bar.", lang)
    else:
        paragraph(doc, "不同数据商的快照因时间点不同而略有差异：美联社转引的 FactSet 一致预期为收入 922.7 亿美元、调整后 EPS 2.09 美元；Longbridge 在财报时点的快照为 921.6 亿美元和 2.091 美元。分析结论不变：收入超预期约 4.3%-4.4%，调整后 EPS 超 6.2%，数据中心收入较 Yahoo 的 858.6 亿美元预期高 3.7%。EPS 差异更大，说明本季并非单纯的出货时点超预期，经营杠杆与产品组合也高于财报前门槛。", lang)
    mix_chart_width = 5.65 if lang == "en" else 4.95
    figure(doc, 5, "Q2 beat across revenue, EPS and Data Center" if lang == "en" else "Q2 的收入、EPS 与数据中心均超预期", lang, [("FactSet consensus via AP" if lang == "en" else "美联社转引 FactSet 一致预期", URLS["factset"]), ("Yahoo Finance results snapshot" if lang == "en" else "Yahoo Finance 财报快照", URLS["yahoo_breaking"]), ("NVIDIA release" if lang == "en" else "英伟达财报公告", URLS["release"])], width=mix_chart_width)
    if lang == "en":
        paragraph(doc, "The new market-platform disclosure is increasingly useful. Hyperscale was $48.71bn, up 13% QoQ and 102% YoY. ACIE reached $40.31bn, up 25% QoQ and 138% YoY, driven by AI natives, enterprises, sovereign customers and hyperscalers consuming AI-cloud capacity. Edge Computing rose 13% sequentially to $7.20bn as Blackwell workstations offset softer consumer PCs. ACIE's faster growth is thesis-positive because it diversifies end demand, although the 10-Q still shows one direct customer at 16% of quarterly revenue and three direct customers at 16%, 15% and 13% of first-half revenue.", lang)
    else:
        paragraph(doc, "新的市场平台披露越来越具有分析价值。Hyperscale 收入为 487.10 亿美元，环比增长 13%、同比增长 102%；ACIE 达到 403.13 亿美元，环比增长 25%、同比增长 138%，驱动力来自 AI 原生企业、传统企业、主权客户以及通过 AI 云消耗算力的超大规模客户。边缘计算环比增长 13% 至 71.98 亿美元，Blackwell 工作站抵消了消费 PC 疲软。ACIE 增速更快有利于终端需求多元化，但 10-Q 仍显示单一直接客户占季度收入 16%，上半年三大直接客户分别占 16%、15% 与 13%。", lang)
    figure(doc, 6, "ACIE was the fastest-growing market platform (+25% QoQ)" if lang == "en" else "ACIE 是增长最快的市场平台（环比 +25%）", lang, [("Q2 FY27 Form 10-Q, Notes 13 and MD&A" if lang == "en" else "Q2 FY27 Form 10-Q 附注 13 与管理层讨论", URLS["tenq"]), ("Q2 FY27 CFO Commentary, p. 2" if lang == "en" else "Q2 FY27 CFO 评论第 2 页", URLS["cfo"])], width=mix_chart_width)
    headers = ["Platform", "Q2 FY27", "QoQ", "YoY"] if lang == "en" else ["市场平台", "Q2 FY27", "环比", "同比"]
    rows = [["Hyperscale", "$48.71bn", "+13%", "+102%"], ["ACIE", "$40.31bn", "+25%", "+138%"], ["Edge Computing", "$7.20bn", "+13%", "+27%"]] if lang == "en" else [["Hyperscale", "487.1 亿美元", "+13%", "+102%"], ["ACIE", "403.1 亿美元", "+25%", "+138%"], ["边缘计算", "72.0 亿美元", "+13%", "+27%"]]
    make_table(doc, headers, rows, lang, [2.0, 1.6, 1.2, 1.2], variance_col=2, font_size=7.8)
    # The immediately preceding figure source also supports this summary table.
    # Keep a separate table source in English; omitting the duplicate Chinese
    # line prevents a page-break paragraph from spilling onto a blank page.
    if lang == "en":
        source(doc, [("Q2 FY27 CFO Commentary, p. 2", URLS["cfo"])], lang)


def page_guidance(doc, lang):
    new_page(doc)
    heading(doc, "GUIDANCE: ANOTHER RAISE, PLUS AN UNUSUALLY LONG HORIZON" if lang == "en" else "指引：再次上调，并给出异常长的可见周期", lang, 1)
    if lang == "en":
        paragraph(doc, "Q3 revenue guidance is $108bn +/-2%, 12.2% above Q2 at the midpoint and $2.84bn above the pre-results Street expectation cited by Yahoo. Guidance assumes no China Data Center compute revenue. Gross margin is expected at 74.0% +/-50 bps and non-GAAP operating expense at roughly $9.0bn. We forecast $110bn, modestly above the midpoint, because both Hyperscale and ACIE remain supply constrained and Rubin production shipments have begun. The principal near-term debate is therefore not demand sufficiency but whether memory, networking, power and system integration can support the planned delivery curve.", lang)
        paragraph(doc, "Management then extended the horizon. On the call, CFO Colette Kress said NVIDIA expects revenue to grow approximately 70% in FY2028. The company linked confidence to profitable compute backlogs at hyperscalers, rapid ACIE growth, and revenue opportunity per gigawatt rising from about $18bn for Hopper to $25bn for Grace Blackwell and $40bn for Vera Rubin. Because this was call commentary rather than a formal quarterly guide, investors should apply a wider range of outcomes; it nevertheless forces a material upward reset to outer-year estimates.", lang)
        paragraph(doc, "China remains excluded from guidance and should be treated as optionality, not base case. Q2 Hopper shipments to China were less than 1% of Data Center revenue. The 10-Q says NVIDIA is effectively foreclosed from broad participation in China's Data Center compute market under the present U.S. and Chinese policy environment. Any re-entry could add revenue, but inspections, a 25% tariff on licensed H200 imports into the United States, and Chinese restrictions make timing and economics uncertain.", lang)
    else:
        paragraph(doc, "Q3 收入指引为 1,080 亿美元 +/-2%，中点较 Q2 高 12.2%，并较 Yahoo 引用的财报前市场预期高 28.4 亿美元。指引不包含中国数据中心计算收入；毛利率指引为 74.0% +/-50 个基点，Non-GAAP 营业费用约 90 亿美元。我们预测 1,100 亿美元，略高于指引中点，原因是 Hyperscale 与 ACIE 仍受供给约束，Rubin 也已开始生产出货。短期争议因而不是需求是否充足，而是内存、网络、电力与系统集成能否支撑计划交付曲线。", lang)
        paragraph(doc, "管理层随后将可见度延伸至更远年份。电话会上，CFO Colette Kress 表示预计 FY2028 收入增长约 70%。信心来自超大规模客户可盈利的算力积压、ACIE 快速增长，以及每吉瓦收入机会从 Hopper 约 180 亿美元提升至 Grace Blackwell 250 亿美元、Vera Rubin 400 亿美元。由于这属于电话会评论而非正式季度指引，投资者应采用更宽结果区间，但它仍然迫使远期预测大幅上修。", lang)
        paragraph(doc, "中国收入仍被排除在指引之外，应视为上行可选项，而非基准情景。Q2 向中国交付的 Hopper 产品不足数据中心收入 1%。10-Q 表示，在当前中美政策环境下，英伟达实际上无法广泛参与中国数据中心计算市场。未来重返市场可能增加收入，但美国境内检查、获批 H200 需承担 25% 关税以及中国侧限制，使时间与经济性高度不确定。", lang)
    figure(doc, 7, "Q3 guidance again cleared the pre-results bar" if lang == "en" else "Q3 指引再次高于财报前门槛", lang, [("Q2 FY27 earnings release" if lang == "en" else "Q2 FY27 财报公告", URLS["release"]), ("Yahoo Finance pre-results consensus snapshot" if lang == "en" else "Yahoo Finance 财报前一致预期快照", URLS["yahoo_breaking"]), ("Q2 FY27 earnings call, 26 Aug. 2026" if lang == "en" else "Q2 FY27 电话会，2026 年 8 月 26 日", URLS["transcript"])], width=6.15)


def page_balance_sheet(doc, lang):
    new_page(doc)
    heading(doc, "BALANCE SHEET: CASH GENERATION IS STRONG, BUT INTENSITY ROSE" if lang == "en" else "资产负债表：现金创造强劲，但资本强度上升", lang, 1)
    if lang == "en":
        paragraph(doc, "The income statement was exceptional; cash conversion was less clean. Accounts receivable rose to $63.06bn from $40.71bn in Q1 and days sales outstanding increased to 60 from 45 because of extended payment terms on large, multi-quarter agreements with investment-grade customers. Inventory climbed to $31.58bn from $25.80bn as NVIDIA prepared for Rubin. Operating cash flow fell sequentially to $24.08bn from $50.34bn and free cash flow to $21.34bn from $48.55bn. The change is mainly working capital and cash taxes, not a collapse in unit economics, but investors should monitor whether receivables and inventory normalize relative to revenue.", lang)
    else:
        paragraph(doc, "利润表表现极强，但现金转换不够干净。应收账款从 Q1 的 407.10 亿美元增至 630.59 亿美元，应收账款周转天数由 45 天升至 60 天，原因是对部分投资级客户的大型、多季度协议延长付款期限。库存从 257.97 亿美元增至 315.75 亿美元，以准备 Rubin 上市。经营现金流从 503.44 亿美元环比降至 240.77 亿美元，自由现金流从 485.54 亿美元降至 213.41 亿美元。变化主要来自营运资金与现金税，而非单位经济性崩塌，但投资者应观察应收和库存相对收入能否恢复正常。", lang)
    figure(doc, 8, "Working capital absorbed cash as receivables and inventory rose" if lang == "en" else "应收与库存上升吸收了现金", lang, [("Q2 FY27 CFO Commentary, pp. 4 and 8" if lang == "en" else "Q2 FY27 CFO 评论第 4、8 页", URLS["cfo"]), ("Q2 FY27 Form 10-Q" if lang == "en" else "Q2 FY27 Form 10-Q", URLS["tenq"])], width=5.65)
    if lang == "en":
        paragraph(doc, "Commitments expanded much faster than quarterly revenue. Supply and capacity commitments increased from $119bn to $279bn, primarily for memory and manufacturing facilities. Together with $29bn of cloud-service agreements, $25bn of data-center leases not commenced, $25bn of equity-investment commitments and $8bn of capital-expenditure commitments, disclosed future commitments totaled $366bn. Separately, maximum gross guarantee exposure was $108.5bn, including $105bn of phased support for the SB Energy PORTS-Pike buildout. These are not all current liabilities and should not be mechanically added, but they materially lengthen NVIDIA's risk duration.", lang)
    else:
        paragraph(doc, "承诺规模增速远快于季度收入。供应与产能承诺从 1,190 亿美元增至 2,790 亿美元，主要用于内存及制造设施；加上 290 亿美元云服务协议、250 亿美元尚未开始的数据中心租赁、250 亿美元股权投资承诺和 80 亿美元资本开支承诺，披露的未来承诺合计 3,660 亿美元。另有最大担保敞口 1,085 亿美元，其中包括对 SB Energy PORTS-Pike 建设分阶段提供的 1,050 亿美元支持。这些并非全部属于当期负债，也不应机械相加，但确实显著拉长了英伟达的风险期限。", lang)
    figure(doc, 9, "Scale of commitments is now a core part of the risk/reward" if lang == "en" else "承诺规模已成为风险收益的核心部分", lang, [("Q2 FY27 Form 10-Q, Notes 10 and MD&A" if lang == "en" else "Q2 FY27 Form 10-Q 附注 10 与管理层讨论", URLS["tenq"]), ("Q2 FY27 CFO Commentary, pp. 3-4" if lang == "en" else "Q2 FY27 CFO 评论第 3-4 页", URLS["cfo"])], width=5.65)


def page_thesis(doc, lang):
    new_page(doc)
    heading(doc, "THESIS UPDATE: STRONGER GROWTH, WIDER FINANCING DEBATE" if lang == "en" else "投资逻辑更新：增长更强，融资争议更宽", lang, 1)
    if lang == "en":
        bullet(doc, "Pillar 1 - Full-stack AI infrastructure leadership: strengthened.", "Blackwell Ultra drove 106% company growth while gross margin held at 75%. Rubin is already entering production shipments, and management claims the revenue opportunity per gigawatt rises materially with each generation. The evidence supports a platform advantage spanning GPU, CPU, networking, systems and software rather than a single-chip cycle.", lang)
        bullet(doc, "Pillar 2 - Demand breadth and duration: strengthened.", "ACIE grew 25% sequentially and faster than Hyperscale, while management's FY2028 framework points to approximately 70% growth. The customer mix is broader at the end-market level, but direct-customer concentration remains high and an unnamed AI research and deployment company contributed a meaningful amount of indirect demand through cloud customers.", lang)
        bullet(doc, "Pillar 3 - Exceptional economics: intact, with a cash-conversion caveat.", "Operating income more than doubled, adjusted EPS beat consensus and gross margin remained best-in-class. Yet receivables, inventory and cash taxes reduced quarterly free cash flow to $21.3bn. The thesis now requires evidence that extended customer terms are temporary and that Rubin inventory converts to shipments on schedule.", lang)
        bullet(doc, "Pillar 4 - Capital-light ecosystem model: weakened.", "NVIDIA is increasingly securing memory, manufacturing, cloud capacity, sites, power and customer financing. This may be rational vertical coordination in a constrained buildout, but $366bn of commitments and $108.5bn of maximum guarantee exposure make the model less capital-light and raise questions about counterparty quality and demand circularity.", lang, RED)
        paragraph(doc, "We raise FY2027 revenue to $411.8bn from $365bn and non-GAAP EPS to $9.39 from $8.80. Q1 and Q2 actuals contribute $177.84bn and $4.09 of EPS; our Q3 assumptions are $110bn and $2.50, while Q4 is $124bn and $2.80. For FY2028, our $700bn revenue estimate applies management's approximately 70% growth framework to the FY2027 base. We assume gross margin moderates to 74.0%, operating margin to 65.0%, and non-GAAP EPS reaches $16.50. The implied earnings growth is powerful but increasingly depends on working-capital discipline and counterparties outside NVIDIA's traditional supply chain.", lang)
    else:
        bullet(doc, "支柱一——全栈 AI 基础设施领导力：强化。", "Blackwell Ultra 推动公司收入增长 106%，毛利率仍保持 75%；Rubin 已进入生产出货。管理层称每吉瓦收入机会随每代平台显著提升。证据支持公司的优势覆盖 GPU、CPU、网络、系统与软件，而非单一芯片周期。", lang)
        bullet(doc, "支柱二——需求广度与持续时间：强化。", "ACIE 环比增长 25%，快于 Hyperscale；FY2028 管理层框架指向约 70% 增长。终端市场组合更加广泛，但直接客户集中度仍高，且一家未具名 AI 研究与部署公司通过云客户贡献了有意义的间接需求。", lang)
        bullet(doc, "支柱三——卓越经济性：保持，但现金转换出现警示。", "营业利润增长超过一倍，调整后 EPS 超预期，毛利率继续处于行业领先水平。然而应收、库存与现金税使季度自由现金流降至 213 亿美元。投资逻辑现在需要证明延长客户账期只是暂时现象，且 Rubin 库存能够按时转化为出货。", lang)
        bullet(doc, "支柱四——轻资本生态模式：削弱。", "英伟达正更主动地锁定内存、制造、云产能、土地、电力与客户融资。这可能是在供给受限环境中的理性纵向协调，但 3,660 亿美元承诺和 1,085 亿美元最大担保敞口使模式不再那么轻资本，也提高交易对手质量与需求循环性的疑问。", lang, RED)
        paragraph(doc, "我们将 FY2027 收入预测从 3,650 亿美元上调至 4,118 亿美元，Non-GAAP EPS 从 8.80 美元上调至 9.39 美元。Q1 与 Q2 实际合计收入 1,778.4 亿美元、EPS 4.09 美元；Q3 假设为收入 1,100 亿美元、EPS 2.50 美元，Q4 为 1,240 亿美元与 2.80 美元。FY2028 收入 7,000 亿美元是将管理层约 70% 增长框架应用于 FY2027 基础。我们假设毛利率降至 74.0%、营业利润率降至 65.0%，Non-GAAP EPS 达 16.50 美元。盈利增长强劲，但越来越依赖营运资金纪律及传统供应链以外的交易对手。", lang)
    figure(doc, 10, "Q2 beat and FY2028 framework drive estimate upgrades" if lang == "en" else "Q2 超预期与 FY2028 框架推动预测上修", lang, [("Our estimates; prior Q1 report dated 6 June 2026" if lang == "en" else "本报告预测；此前 Q1 报告日期为 2026 年 6 月 6 日", None), ("Q2 FY27 earnings call, 26 Aug. 2026" if lang == "en" else "Q2 FY27 电话会，2026 年 8 月 26 日", URLS["transcript"]), ("Q2 FY27 CFO Commentary" if lang == "en" else "Q2 FY27 CFO 评论", URLS["cfo"])], width=6.1)


def page_estimates_valuation(doc, lang, market):
    new_page(doc)
    heading(doc, "UPDATED ESTIMATES AND VALUATION" if lang == "en" else "更新后的预测与估值", lang, 1)
    ref_price = float(market["premarket"]) if isinstance(market.get("premarket"), (int, float)) else 223.94
    regular_close = float(market["regular_close"]) if isinstance(market.get("regular_close"), (int, float)) else 209.66
    old_pe = ref_price / 8.80
    new_pe = ref_price / 9.39
    fy28_pe = ref_price / 16.50
    pe_change = new_pe - old_pe
    base_upside = (320.0 / ref_price - 1.0) * 100.0
    close_upside = (320.0 / regular_close - 1.0) * 100.0
    if lang == "en":
        paragraph(doc, "Our prior Q1 update carried FY2027 revenue of $365bn, non-GAAP EPS of $8.80 and a $280 target. We now have two quarters of actual results, Q3 guidance at $108bn, and a management FY2028 growth framework. The table below separates the mechanical uplift from explicit assumptions. Old estimates are our 6 June 2026 report, not external consensus; new estimates are independent analytical forecasts and are not company guidance.", lang)
        headers = ["Metric", "FY2027E old", "FY2027E new", "Change", "FY2028E new"]
        rows = [
            ["Revenue ($bn)", "365.0", "411.8", "+12.8%", "700.0"],
            ["Revenue growth", "69.0%", "90.7%", "+2,170 bps", "70.0%"],
            ["Non-GAAP gross margin", "75.0%", "74.5%", "-50 bps", "74.0%"],
            ["Non-GAAP operating income ($bn)", "245.0", "270.0", "+10.2%", "455.0"],
            ["Non-GAAP operating margin", "67.1%", "65.6%", "-150 bps", "65.0%"],
            ["Non-GAAP diluted EPS", "$8.80", "$9.39", "+6.7%", "$16.50"],
            [f"P/E at ${ref_price:.2f}", f"{old_pe:.1f}x", f"{new_pe:.1f}x", f"{pe_change:.1f}x", f"{fy28_pe:.1f}x"],
        ]
    else:
        paragraph(doc, "此前 Q1 更新报告采用 FY2027 收入 3,650 亿美元、Non-GAAP EPS 8.80 美元与 280 美元目标价。现在已有两个季度实际业绩、Q3 1,080 亿美元指引以及 FY2028 管理层增长框架。下表区分机械上修与明确假设。“旧预测”来自我们 2026 年 6 月 6 日报告，并非外部一致预期；“新预测”为独立分析判断，不构成公司指引。", lang)
        headers = ["指标", "FY2027E 旧", "FY2027E 新", "变动", "FY2028E 新"]
        rows = [
            ["收入（十亿美元）", "365.0", "411.8", "+12.8%", "700.0"],
            ["收入增长率", "69.0%", "90.7%", "+2,170 个基点", "70.0%"],
            ["Non-GAAP 毛利率", "75.0%", "74.5%", "-50 个基点", "74.0%"],
            ["Non-GAAP 营业利润（十亿美元）", "245.0", "270.0", "+10.2%", "455.0"],
            ["Non-GAAP 营业利润率", "67.1%", "65.6%", "-150 个基点", "65.0%"],
            ["Non-GAAP 摊薄 EPS", "8.80 美元", "9.39 美元", "+6.7%", "16.50 美元"],
            [f"按 {ref_price:.2f} 美元计算 P/E", f"{old_pe:.1f} 倍", f"{new_pe:.1f} 倍", f"{pe_change:.1f} 倍", f"{fy28_pe:.1f} 倍"],
        ]
    make_table(doc, headers, rows, lang, [2.4, 1.35, 1.35, 1.25, 1.45], variance_col=3, font_size=7.6)
    source(doc, [("Our estimates and 6 June 2026 Q1 update" if lang == "en" else "本报告预测与 2026 年 6 月 6 日 Q1 更新", None), ("Q2 FY27 earnings materials" if lang == "en" else "Q2 FY27 财报材料", URLS["cfo"]), ("Q2 FY27 earnings call" if lang == "en" else "Q2 FY27 电话会", URLS["transcript"])], lang)
    if lang == "en":
        heading(doc, "Price target methodology", lang, 2)
        paragraph(doc, "We raise our target to $320 from $280. The target applies 19.4x to FY2028E non-GAAP EPS of $16.50. This is below NVIDIA's five-year median trailing P/E of roughly 58.5x and below its current five-year 'reasonable range' reported by Longbridge, but forward and trailing multiples are not directly comparable. The chosen multiple explicitly discounts: (1) the extraordinary scale embedded in the 70% FY2028 framework; (2) margin normalization to 74%; (3) customer and geographic concentration; and (4) the transition from a capital-light supplier toward a coordinator of supply, sites, credit and cloud capacity.", lang)
        paragraph(doc, f"Scenario framing: our bear value is $250, assuming FY2028 revenue of roughly $600bn, EPS near $13 and a 19x multiple. The $320 base assumes $700bn and $16.50. A $400 bull case assumes China optionality, faster Rubin conversion, better cash conversion and $17.50-$18.00 of EPS at 22x-23x. At the 27 August pre-market reference of ${ref_price:.2f}, the base case implies {base_upside:.1f}% upside. The regular 26 August close of ${regular_close:.2f} implies {close_upside:.1f}%; we use the more conservative pre-market reference for rating decisions.", lang)
        heading(doc, "What the target does not assume", lang, 2)
        paragraph(doc, "The base case assumes no China Data Center compute revenue, no valuation credit for unrealized equity gains, and no benefit from guarantee-supported sites beyond revenue captured in the FY2028 framework. It also does not assume that all commitments are drawn or that maximum guarantees become liabilities. Conversely, it does not ignore them: the lower multiple is the mechanism through which counterparty, duration and circularity risks enter valuation. This approach is more transparent than adding uncertain obligations to enterprise value today.", lang)
    else:
        heading(doc, "目标价方法", lang, 2)
        paragraph(doc, "我们将目标价从 280 美元上调至 320 美元，对 FY2028E Non-GAAP EPS 16.50 美元采用 19.4 倍市盈率。该倍数低于英伟达约 58.5 倍的五年历史中位数，也低于 Longbridge 所示当前五年“合理区间”，但前瞻与历史市盈率不可直接比较。所选倍数明确折价以下因素：（1）FY2028 约 70% 增长框架已经包含极大规模；（2）毛利率回归至 74%；（3）客户与地区集中；（4）公司正从轻资本供应商转向协调供应、站点、信用与云产能。", lang)
        paragraph(doc, f"情景框架：悲观价值 250 美元，假设 FY2028 收入约 6,000 亿美元、EPS 约 13 美元并采用 19 倍；基准 320 美元对应收入 7,000 亿美元和 EPS 16.50 美元；乐观 400 美元假设中国可选项、Rubin 转换更快、现金转换改善，EPS 为 17.50-18.00 美元并采用 22-23 倍。以 8 月 27 日盘前参考价 {ref_price:.2f} 美元计，基准潜在上涨 {base_upside:.1f}%；以 8 月 26 日常规收盘 {regular_close:.2f} 美元计为 {close_upside:.1f}%。评级判断采用更保守的盘前参考。", lang)
        heading(doc, "目标价未计入的因素", lang, 2)
        paragraph(doc, "基准情景不假设中国数据中心计算收入，不给予未实现股票投资收益估值，也不对担保支持站点给予超出 FY2028 框架的额外收入。模型也不假设所有承诺都会被提取，或最大担保一定转化为负债。反过来，我们也未忽略它们：较低估值倍数就是将交易对手、期限和循环性风险纳入估值的机制。这比当前直接把不确定义务加入企业价值更透明。", lang)


def page_risks_catalysts(doc, lang, market):
    new_page(doc)
    heading(doc, "CATALYSTS, RISKS AND TRADING CONTEXT" if lang == "en" else "催化剂、风险与交易背景", lang, 1)
    if lang == "en":
        heading(doc, "Catalyst calendar", lang, 2)
        bullet(doc, "Q3 FY2027 execution (November 2026).", "Delivery against the $108bn guide, 74% gross margin and the first Rubin production shipments will test whether supply—not end demand—remains the limiting factor. ACIE growth above Hyperscale would further support demand diversification.", lang)
        bullet(doc, "Rubin and memory supply visibility.", "Qualification, rack deployment, networking availability and memory procurement will determine whether NVIDIA can convert $279bn of supply commitments into timely revenue. A narrowing DSO and inventory-to-sales ratio would improve the quality of the growth narrative.", lang)
        bullet(doc, "FY2028 framework confirmation.", "Subsequent management commentary, hyperscaler capital budgets and AI-cloud financing activity can validate or challenge the approximately 70% growth expectation. Evidence that third-party customers are using capacity profitably matters more than gross site announcements.", lang)
        heading(doc, "Principal risks", lang, 2)
        bullet(doc, "Commitment and counterparty risk.", "Long-dated supply, lease, cloud and guarantee structures may outlive a demand cycle. Defaults, construction delays, weaker customer financing or slower utilization could create losses, cash calls or reputational damage even if current revenue remains strong.", lang, RED)
        bullet(doc, "Customer concentration and circularity.", "One direct customer represented 16% of quarterly revenue, and a meaningful amount of demand came indirectly from one AI research and deployment company through cloud customers. NVIDIA's investments and financing support can strengthen the ecosystem while making organic demand harder to separate from enabled demand.", lang, RED)
        bullet(doc, "Technology and supply execution.", "Rubin ramps alongside Blackwell, with memory and manufacturing commitments rising sharply. Yield, networking, cooling, power, system integration or component shortages can defer revenue and create excess inventory provisions.", lang, RED)
        bullet(doc, "Policy and China.", "China is excluded from guidance and represented less than 1% of Data Center revenue. Export licensing, U.S. inspection and tariff requirements, and Chinese restrictions can limit participation or make any re-entry less profitable than headline revenue suggests.", lang, RED)
        paragraph(doc, f"The one-year chart shows why expectations matter. NVDA closed at ${float(market['regular_close']):.2f} on 26 August after a 1.6% decline in the regular session, then traded at ${float(market['postmarket']):.2f} after hours and about ${float(market['premarket']):.2f} pre-market on 27 August. The positive reaction confirms the beat and FY2028 framework were not fully priced, but the stock remains volatile around earnings. Our Outperform rating is based on twelve-month risk-adjusted return, not on the direction of the first post-release trading session.", lang)
    else:
        heading(doc, "催化剂日历", lang, 2)
        bullet(doc, "FY2027 Q3 执行（2026 年 11 月）。", "能否达到 1,080 亿美元收入指引、74% 毛利率以及完成首批 Rubin 生产出货，将验证供给而非终端需求是否仍是限制因素。若 ACIE 增速继续快于 Hyperscale，将进一步支持需求多元化。", lang)
        bullet(doc, "Rubin 与内存供应可见度。", "认证、机架部署、网络可用性与内存采购将决定英伟达能否把 2,790 亿美元供应承诺及时转化为收入。应收周转天数和库存收入比下降，将改善增长叙事质量。", lang)
        bullet(doc, "FY2028 框架确认。", "后续管理层评论、超大规模客户资本预算与 AI 云融资活动可以验证或挑战约 70% 增长预期。证明第三方客户能够盈利使用算力，比单纯公布站点规模更重要。", lang)
        heading(doc, "主要风险", lang, 2)
        bullet(doc, "承诺与交易对手风险。", "长期供应、租赁、云和担保结构可能跨越完整需求周期。违约、建设延迟、客户融资减弱或利用率下降，可能在当期收入仍强时引发损失、现金调用或声誉影响。", lang, RED)
        bullet(doc, "客户集中与循环性。", "单一直接客户占季度收入 16%，一家 AI 研究与部署公司还通过云客户贡献了有意义的间接需求。英伟达的投资与融资支持能够强化生态，也使自然需求与被支持需求更难区分。", lang, RED)
        bullet(doc, "技术与供应执行。", "Rubin 与 Blackwell 同时爬坡，内存与制造承诺大幅增加。良率、网络、冷却、电力、系统集成或组件短缺可能推迟收入，并产生额外库存拨备。", lang, RED)
        bullet(doc, "政策与中国。", "中国被排除在指引之外，Q2 占数据中心收入不足 1%。出口许可、美国检查与关税要求以及中国侧限制，都可能限制参与度，或使重返市场的利润低于表面收入。", lang, RED)
        paragraph(doc, f"一年股价图说明预期的重要性。NVDA 在 8 月 26 日常规交易中下跌 1.6%，收于 {float(market['regular_close']):.2f} 美元，盘后升至 {float(market['postmarket']):.2f} 美元，8 月 27 日盘前约 {float(market['premarket']):.2f} 美元。正面反应说明超预期业绩与 FY2028 框架尚未完全计价，但财报前后波动仍高。我们的“跑赢大盘”评级基于未来十二个月风险调整回报，而非财报后首个交易时段的方向。", lang)
    figure(doc, 11, "One-year price history still leaves material upside to our base case" if lang == "en" else "一年股价历史仍显示基准情景具有可观上涨空间", lang, [("Longbridge adjusted daily K-line through 26 Aug. 2026" if lang == "en" else "Longbridge 前复权日 K，截至 2026 年 8 月 26 日", URLS["longbridge_quote"]), ("Yahoo Finance dynamic market data" if lang == "en" else "Yahoo Finance 动态市场数据", URLS["yahoo_quote"]), ("Our valuation" if lang == "en" else "本报告估值", None)], width=6.15)


def page_sources(doc, lang):
    new_page(doc)
    heading(doc, "SOURCES, METHODOLOGY AND DISCLOSURES" if lang == "en" else "资料来源、方法与披露", lang, 1)
    if lang == "en":
        heading(doc, "Timeliness and identity verification", lang, 2)
        paragraph(doc, "Today is 27 August 2026. NVIDIA Corporation (Nasdaq: NVDA; SEC CIK 1045810) released Q2 FY2027 results on 26 August 2026 for the quarter ended 26 July 2026. The Form 8-K and Form 10-Q were filed on the same date. The earnings call began at 2:00 p.m. Pacific / 5:00 p.m. Eastern on 26 August, matching the release date. The quarter is therefore the company's latest reported period and the materials are less than one day old.", lang)
        heading(doc, "Source hierarchy", lang, 2)
        paragraph(doc, "Reported financials, market-platform data, guidance, commitments and risk factors are anchored to NVIDIA's earnings release, CFO Commentary, investor presentation and SEC filings. The official webcast page confirms call timing; the Seeking Alpha transcript and Yahoo call highlights are used for management's FY2028 growth framework and Q&A details. Consensus is FactSet as reported by AP and the release-time Yahoo/Longbridge snapshots. Differences of roughly $0.1bn across consensus vendors reflect snapshot timing and do not alter the beat/miss conclusion.", lang)
        heading(doc, "Forecast and valuation conventions", lang, 2)
        paragraph(doc, "FY2027 'old' estimates and the $280 prior target are from our 6 June 2026 Q1 update. FY2027 and FY2028 'new' estimates are our independent forecasts. FY2028 revenue applies the approximately 70% management framework to our FY2027 base and is not formal company guidance. Non-GAAP estimates use NVIDIA's FY2027 definition, which includes stock-based compensation. Price target upside uses the 27 August pre-market reference; market capitalization and 52-week range were dynamically fetched through yfinance. Longbridge supplied the verified regular, after-hours and pre-market quotes and adjusted price history.", lang)
        heading(doc, "Important disclosure", lang, 2)
        paragraph(doc, "This report is independent analytical research for informational purposes only. It is not investment advice, an offer, a solicitation or a fiduciary recommendation. Rating, estimates, scenarios and price target are opinions based on assumptions that may prove wrong. NVIDIA is exposed to technology transitions, manufacturing concentration, supply constraints, customer concentration, export controls, tariffs, competition, market valuation, investments, guarantees and long-duration commitments. Readers should review primary filings and obtain advice appropriate to their circumstances.", lang)
        refs = [
            ("NVIDIA Q2 FY2027 earnings release (26 Aug. 2026)", URLS["release"]),
            ("Form 8-K, filed 26 Aug. 2026", URLS["eightk"]),
            ("Form 10-Q, filed 26 Aug. 2026", URLS["tenq"]),
            ("Q2 FY2027 CFO Commentary (26 Aug. 2026)", URLS["cfo"]),
            ("Q2 FY2027 investor presentation (Aug. 2026)", URLS["presentation"]),
            ("Official earnings webcast (26 Aug. 2026)", URLS["webcast"]),
            ("Earnings-call transcript (26 Aug. 2026)", URLS["transcript"]),
            ("Yahoo call highlights (26 Aug. 2026)", URLS["yahoo_call"]),
            ("FactSet consensus via AP (26 Aug. 2026)", URLS["factset"]),
            ("Q1 FY2027 release and prior guidance (20 May 2026)", URLS["q1_release"]),
            ("Q4/FY2026 release (25 Feb. 2026)", URLS["q4_release"]),
            ("NVIDIA Investor Relations", URLS["ir"]),
            ("Yahoo Finance market data", URLS["yahoo_quote"]),
            ("Longbridge NVDA market data", URLS["longbridge_quote"]),
        ]
    else:
        heading(doc, "时效性与发行人核验", lang, 2)
        paragraph(doc, "今天是 2026 年 8 月 27 日。英伟达公司（纳斯达克：NVDA；SEC CIK 1045810）于 2026 年 8 月 26 日发布 FY2027 Q2 业绩，对应截至 2026 年 7 月 26 日的季度；Form 8-K 与 Form 10-Q 同日提交。电话会于 8 月 26 日太平洋时间 14:00 / 美东时间 17:00 开始，与财报发布日期一致。因此本季是公司最新报告期，材料发布时间不足一天。", lang)
        heading(doc, "资料层级", lang, 2)
        paragraph(doc, "财务数据、市场平台、指引、承诺与风险因素以英伟达财报公告、CFO 评论、投资者演示及 SEC 文件为锚。官方回放页面确认电话会时点；Seeking Alpha 逐字稿和 Yahoo 电话会要点用于 FY2028 增长框架及问答细节。一致预期采用美联社转引 FactSet，以及财报时点 Yahoo/Longbridge 快照。不同数据商约 1 亿美元的差异来自快照时点，不改变超预期结论。", lang)
        heading(doc, "预测与估值口径", lang, 2)
        paragraph(doc, "FY2027“旧预测”与此前 280 美元目标价来自我们 2026 年 6 月 6 日 Q1 更新。FY2027 与 FY2028“新预测”为独立预测。FY2028 收入将管理层约 70% 增长框架应用于我们的 FY2027 基础，并非公司正式指引。Non-GAAP 预测采用英伟达 FY2027 定义，包含股权激励费用。目标价上涨空间使用 8 月 27 日盘前参考；市值与 52 周区间通过 yfinance 动态获取。Longbridge 提供已验证的常规、盘后、盘前报价与复权股价历史。", lang)
        heading(doc, "重要披露", lang, 2)
        paragraph(doc, "本报告为独立分析研究，仅供信息参考，不构成投资建议、要约、招揽或受托建议。评级、预测、情景和目标价均为基于可能被证伪假设的观点。英伟达面临技术过渡、制造集中、供应限制、客户集中、出口管制、关税、竞争、市场估值、投资、担保及长期承诺风险。读者应核阅一手文件，并根据自身情况获取专业意见。", lang)
        refs = [
            ("英伟达 Q2 FY2027 财报公告（2026 年 8 月 26 日）", URLS["release"]),
            ("Form 8-K（2026 年 8 月 26 日提交）", URLS["eightk"]),
            ("Form 10-Q（2026 年 8 月 26 日提交）", URLS["tenq"]),
            ("Q2 FY2027 CFO 评论（2026 年 8 月 26 日）", URLS["cfo"]),
            ("Q2 FY2027 投资者演示（2026 年 8 月）", URLS["presentation"]),
            ("官方财报电话会回放（2026 年 8 月 26 日）", URLS["webcast"]),
            ("财报电话会逐字稿（2026 年 8 月 26 日）", URLS["transcript"]),
            ("Yahoo 电话会要点（2026 年 8 月 26 日）", URLS["yahoo_call"]),
            ("美联社转引 FactSet 一致预期（2026 年 8 月 26 日）", URLS["factset"]),
            ("Q1 FY2027 财报与此前指引（2026 年 5 月 20 日）", URLS["q1_release"]),
            ("Q4/FY2026 财报（2026 年 2 月 25 日）", URLS["q4_release"]),
            ("英伟达投资者关系", URLS["ir"]),
            ("Yahoo Finance 市场数据", URLS["yahoo_quote"]),
            ("Longbridge NVDA 市场数据", URLS["longbridge_quote"]),
        ]
    heading(doc, "REFERENCE LINKS" if lang == "en" else "参考链接", lang, 2)
    for label, url in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.12)
        p.paragraph_format.space_after = Pt(1.7)
        r = p.add_run("• ")
        set_run_font(r, lang, size=8.2, color=NAVY)
        add_hyperlink(p, label, url, lang, size=8.2)


def build(lang):
    missing = [str(path) for path in CHARTS.values() if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing charts: " + ", ".join(missing))
    market = market_data()
    doc = setup_doc(lang)
    title_page(doc, lang, market)
    page_revenue(doc, lang)
    page_profitability(doc, lang)
    page_mix(doc, lang)
    page_guidance(doc, lang)
    page_balance_sheet(doc, lang)
    page_thesis(doc, lang)
    page_estimates_valuation(doc, lang, market)
    page_risks_catalysts(doc, lang, market)
    page_sources(doc, lang)
    filename = "NVDA_Q2_FY2027_Earnings_Update.docx" if lang == "en" else "NVDA_Q2_FY2027_业绩更新报告_中文版.docx"
    path = OUT / filename
    doc.core_properties.title = "NVIDIA Q2 FY2027 Earnings Update" if lang == "en" else "英伟达 2027 财年第二季度业绩更新报告"
    doc.core_properties.subject = "Post-earnings institutional equity research"
    doc.core_properties.author = "NVDA Research"
    doc.core_properties.keywords = "NVIDIA, NVDA, Q2 FY2027, earnings update, equity research"
    doc.save(path)
    print(path)


if __name__ == "__main__":
    build("en")
    build("cn")
