#!/usr/bin/env python3
"""Deterministic structural QA for the WDC Q1 FY2027 bilingual reports."""

from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[2]
DOCS = {
    "en": ROOT / "output/WDC/WDC_Q1_FY2027_Earnings_Update.docx",
    "cn": ROOT / "output/WDC/WDC_Q1_FY2027_业绩更新报告_中文版.docx",
}
PLACEHOLDER_RE = re.compile(r"\b(?:TODO|TBD|PLACEHOLDER|XXX)\b", re.I)
URL_RE = re.compile(r"https?://")
FIGURE_RE = re.compile(r"^(?:Figure|图)\s*(\d+)[.：:]")


def audit(path: Path, lang: str) -> dict:
    assert path.exists() and path.stat().st_size > 100_000, path
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
        document_xml = archive.read("word/document.xml").decode("utf-8")
        rels_xml = archive.read("word/_rels/document.xml.rels").decode("utf-8")
        package_xml = "".join(
            archive.read(name).decode("utf-8", errors="ignore")
            for name in archive.namelist()
            if name.endswith(".xml")
        )

    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    captions = [p.text for p in doc.paragraphs if FIGURE_RE.match(p.text)]
    figures = [FIGURE_RE.match(p).group(1) for p in captions]
    hyperlinks = [
        rel.target_ref
        for rel in doc.part.rels.values()
        if rel.reltype.endswith("/hyperlink")
    ]
    descriptions = []
    for shape in doc.inline_shapes:
        doc_pr = shape._inline.docPr
        descriptions.append(doc_pr.get("descr") or "")

    tables = doc.tables
    for table in tables:
        tbl_pr = table._tbl.tblPr
        assert tbl_pr.find(qn("w:tblW")) is not None
        assert tbl_pr.find(qn("w:tblInd")) is not None
        assert table._tbl.tblGrid is not None
        first_row_pr = table.rows[0]._tr.get_or_add_trPr()
        assert first_row_pr.find(qn("w:tblHeader")) is not None
        for row in table.rows:
            for cell in row.cells:
                assert cell._tc.tcPr.find(qn("w:tcW")) is not None

    assert not PLACEHOLDER_RE.search(text)
    assert not URL_RE.search(text)
    assert len(tables) == 3
    assert len(doc.inline_shapes) == 11
    assert sorted(figures, key=int) == [str(i) for i in range(1, 12)]
    assert len(set(captions)) == 11
    assert descriptions and all(descriptions) and len(set(descriptions)) == 11
    hyperlink_instances = document_xml.count("<w:hyperlink")
    assert len(hyperlinks) >= 12 and hyperlink_instances >= 25
    assert all(url.startswith("https://") for url in hyperlinks)
    assert len(headings) >= 9
    assert document_xml.count('w:type="page"') == 9
    assert "TargetMode=\"External\"" in rels_xml
    assert "PAGE" in package_xml

    if lang == "en":
        word_count = len(re.findall(r"\b[A-Za-z0-9][A-Za-z0-9'’.-]*\b", text))
        assert 3_000 <= word_count <= 5_000
        language_metric = word_count
    else:
        chinese_chars = len(re.findall(r"[\u3400-\u9fff]", text))
        assert chinese_chars >= 4_000
        language_metric = chinese_chars

    return {
        "file": str(path),
        "size_bytes": path.stat().st_size,
        "tables": len(tables),
        "figures": len(doc.inline_shapes),
        "captions": len(captions),
        "hyperlink_instances": hyperlink_instances,
        "unique_hyperlink_targets": len(hyperlinks),
        "headings": len(headings),
        "manual_page_breaks": document_xml.count('w:type="page"'),
        "word_count" if lang == "en" else "chinese_characters": language_metric,
        "zip_integrity": "pass",
        "placeholders": 0,
        "raw_urls": 0,
    }


def main() -> None:
    results = {lang: audit(path, lang) for lang, path in DOCS.items()}
    print(json.dumps(results, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
