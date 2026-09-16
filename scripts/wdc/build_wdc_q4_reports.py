#!/usr/bin/env python3
"""Build English and complete Chinese Western Digital Q4 FY2026 earnings updates."""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/macrossz/DevTools/VscodeProject/ClaudeCode/financial_analysis")
OUT = ROOT / "output" / "WDC"
DATA = ROOT / "data" / "wdc"
OUT.mkdir(parents=True, exist_ok=True)
DATA.mkdir(parents=True, exist_ok=True)

DOC_SKILL_SCRIPTS = Path(
    "/Users/macrossz/.codex/plugins/cache/openai-primary-runtime/documents/"
    "26.905.11957/skills/documents/scripts"
)
sys.path.insert(0, str(DOC_SKILL_SCRIPTS))
from table_geometry import apply_table_geometry, column_widths_from_weights  # noqa: E402


PAGE_WIDTH_DXA = 10166
NAVY = "123B5D"
BLUE = "2F6B9A"
TEAL = "008A8A"
GREEN = "2F7D32"
ORANGE = "B85C00"
RED = "B42318"
GREY = "667085"
LIGHT = "F2F4F7"
WHITE = "FFFFFF"
BLACK = "000000"

URLS = {
    "release": "https://investor.wdc.com/node/28586/pdf",
    "presentation": "https://investor.wdc.com/static-files/e1f02f77-4432-42c3-8bd4-024371f40e54",
    "webcast": "https://investor.wdc.com/events/event-details/western-digital-fourth-quarter-fiscal-2026-earnings-call",
    "transcript": "https://seekingalpha.com/article/4931243-western-digital-corporation-wdc-q4-2026-earnings-call-transcript",
    "transcript_full": "https://earningscalls.dev/transcripts/western-digital-corporation_wdc_earnings_call_transcript_2026-08-05",
    "eightk": "https://www.sec.gov/Archives/edgar/data/106040/000162828026053305/0001628280-26-053305-index.htm",
    "tenk": "https://www.sec.gov/Archives/edgar/data/106040/000162828026057139/0001628280-26-057139-index.html",
    "tenk_doc": "https://www.sec.gov/Archives/edgar/data/106040/000162828026057139/wdc-20260703.htm",
    "q3_release": "https://www.westerndigital.com/company/newsroom/press-releases/2026/2026-04-30-wd-reports-fiscal-third-quarter-2026-financial-results",
    "q3_presentation": "https://investor.wdc.com/static-files/5b2d41c1-7d45-4575-b9ea-c51424dbffeb",
    "consensus": "https://www.marketbeat.com/stocks/NASDAQ/WDC/earnings/",
    "yahoo": "https://finance.yahoo.com/quote/WDC/",
    "longbridge": "https://longbridge.com/en/quote/WDC.US",
    "ir": "https://investor.wdc.com/",
}

CHARTS = {
    1: OUT / "wdc_q4_chart1_quarterly_revenue.png",
    2: OUT / "wdc_q4_chart2_eps.png",
    3: OUT / "wdc_q4_chart3_margins.png",
    4: OUT / "wdc_q4_chart4_free_cash_flow.png",
    5: OUT / "wdc_q4_chart5_end_market_mix.png",
    6: OUT / "wdc_q4_chart6_exabytes.png",
    7: OUT / "wdc_q4_chart7_prior_guidance.png",
    8: OUT / "wdc_q4_chart8_price_cost.png",
    9: OUT / "wdc_q4_chart9_estimate_revisions.png",
    10: OUT / "wdc_q4_chart10_valuation_scenarios.png",
    11: OUT / "wdc_q4_chart11_price_history.png",
}


def dynamic_market_data() -> dict:
    """Fetch yfinance fields dynamically with Python 3.13 and cross-check Longbridge."""
    result = {
        "price": "N/A",
        "market_cap": "N/A",
        "high": "N/A",
        "low": "N/A",
        "previous_close": "N/A",
        "longbridge_last": "N/A",
    }
    code = """
import json, yfinance as yf
info = yf.Ticker('WDC').fast_info
print(json.dumps({
    'price': round(float(info.last_price), 2),
    'market_cap': float(info.market_cap),
    'high': round(float(info.year_high), 2),
    'low': round(float(info.year_low), 2),
    'previous_close': round(float(info.previous_close), 2),
}))
"""
    try:
        proc = subprocess.run(
            ["python3.13", "-c", code],
            check=True,
            capture_output=True,
            text=True,
            timeout=30,
        )
        result.update(json.loads(proc.stdout.strip()))
    except Exception as exc:
        print(f"WARNING: yfinance market-data retrieval failed for WDC: {exc}")
    quote_path = DATA / "longbridge_quote.json"
    if quote_path.exists():
        try:
            quote = json.loads(quote_path.read_text())[0]
            result["longbridge_last"] = float(quote["last"])
        except Exception as exc:
            print(f"WARNING: Longbridge quote parsing failed for WDC: {exc}")
    return result


def set_run_font(run, lang: str, size=9.5, bold=False, italic=False, color=BLACK, heading=False):
    if lang == "cn":
        face = "Heiti SC" if heading else "Songti SC"
    else:
        face = "Times New Roman"
    run.font.name = face
    run._element.get_or_add_rPr()
    run._element.rPr.rFonts.set(qn("w:ascii"), face)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), face)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), face)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    if lang == "cn":
        language = run._element.rPr.find(qn("w:lang"))
        if language is None:
            language = OxmlElement("w:lang")
            run._element.rPr.append(language)
        language.set(qn("w:val"), "zh-CN")
        language.set(qn("w:eastAsia"), "zh-CN")


def paragraph(doc, text, lang, size=9.5, bold=False, italic=False, color=BLACK, before=0, after=4.0, align=None, keep=False, heading=False, style=None):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    set_run_font(r, lang, size=size, bold=bold, italic=italic, color=color, heading=heading)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.04
    p.paragraph_format.keep_with_next = keep
    if align is not None:
        p.alignment = align
    return p


def heading(doc, text, lang, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(1.5 if level == 1 else 0.5)
    p.paragraph_format.space_after = Pt(4.2)
    p.paragraph_format.keep_with_next = True
    return p


def bullet(doc, title, body, lang, color=NAVY, compact=False):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.03)
    p.paragraph_format.first_line_indent = Inches(-0.03)
    p.paragraph_format.space_after = Pt(3.0 if compact else 4.0)
    p.paragraph_format.line_spacing = 1.03
    r = p.add_run("■ ")
    set_run_font(r, lang, size=8.8 if compact else 9.25, bold=True, color=color, heading=True)
    r = p.add_run(title + " ")
    set_run_font(r, lang, size=8.8 if compact else 9.25, bold=True, color=BLACK, heading=True)
    r = p.add_run(body)
    set_run_font(r, lang, size=8.8 if compact else 9.25)
    return p


def add_hyperlink(paragraph_obj, label, url, lang, size=7.4):
    rel_id = paragraph_obj.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    face = "Songti SC" if lang == "cn" else "Times New Roman"
    for slot in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        fonts.set(qn(slot), face)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(round(size * 2)))
    for element in (fonts, color, underline, sz):
        props.append(element)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph_obj._p.append(link)


def source(doc, items, lang, lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0.5)
    p.paragraph_format.space_after = Pt(3.0)
    p.paragraph_format.line_spacing = 1.0
    prefix = lead or ("Sources: " if lang == "en" else "资料来源：")
    r = p.add_run(prefix)
    set_run_font(r, lang, size=7.4, italic=True, color=GREY)
    for idx, (label, url) in enumerate(items):
        if idx:
            r = p.add_run("; ")
            set_run_font(r, lang, size=7.4, italic=True, color=GREY)
        if url:
            add_hyperlink(p, label, url, lang, size=7.4)
        else:
            r = p.add_run(label)
            set_run_font(r, lang, size=7.4, italic=True, color=GREY)
    r = p.add_run(".")
    set_run_font(r, lang, size=7.4, italic=True, color=GREY)
    return p


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def make_table(doc, headers, rows, lang, weights, variance_col=None, font_size=7.8):
    table_obj = doc.add_table(rows=1, cols=len(headers))
    table_obj.style = "Table Grid"
    table_obj.alignment = WD_TABLE_ALIGNMENT.LEFT
    for col, label in enumerate(headers):
        cell = table_obj.rows[0].cells[col]
        cell.text = label
        shade(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cell.paragraphs[0].runs:
            set_run_font(run, lang, size=font_size, bold=True, color=WHITE, heading=True)
    repeat_header(table_obj.rows[0])
    for row_idx, values in enumerate(rows):
        cells = table_obj.add_row().cells
        for col, value in enumerate(values):
            cells[col].text = str(value)
            set_cell_margins(cells[col])
            if row_idx % 2:
                shade(cells[col], LIGHT)
            cells[col].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[col].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if col == 0 else WD_ALIGN_PARAGRAPH.CENTER
            for run in cells[col].paragraphs[0].runs:
                color = BLACK
                if variance_col == col:
                    color = GREEN if str(value).startswith("+") else RED if str(value).startswith("-") else BLACK
                set_run_font(run, lang, size=font_size, bold=(col == 0), color=color)
    widths = column_widths_from_weights(weights, PAGE_WIDTH_DXA)
    apply_table_geometry(table_obj, widths, table_width_dxa=PAGE_WIDTH_DXA, indent_dxa=120)
    return table_obj


def figure(doc, idx, caption, lang, sources, width=5.65):
    label = f"Figure {idx}. {caption}" if lang == "en" else f"图 {idx}：{caption}"
    paragraph(doc, label, lang, size=8.25, bold=True, color=BLACK, after=1.0, keep=True, heading=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0.3)
    shape = p.add_run().add_picture(str(CHARTS[idx]), width=Inches(width))
    shape._inline.docPr.set("descr", caption)
    source(doc, sources, lang)


def page_field(paragraph_obj):
    run = paragraph_obj.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])
    set_run_font(run, "en", size=7.6, color=GREY)


def setup_doc(lang):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.64)
    section.bottom_margin = Inches(0.60)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.27)
    section.footer_distance = Inches(0.29)

    normal = doc.styles["Normal"]
    normal_face = "Songti SC" if lang == "cn" else "Times New Roman"
    normal.font.name = normal_face
    for slot in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        normal._element.rPr.rFonts.set(qn(slot), normal_face)
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(4.0)
    normal.paragraph_format.line_spacing = 1.04

    title = doc.styles["Title"]
    title_face = "Heiti SC" if lang == "cn" else "Times New Roman"
    title.font.name = title_face
    for slot in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        title._element.rPr.rFonts.set(qn(slot), title_face)
    title.font.size = Pt(17.5)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)

    for style_name, size in (("Heading 1", 13.5), ("Heading 2", 10.8), ("Heading 3", 9.8)):
        style = doc.styles[style_name]
        face = "Heiti SC" if lang == "cn" else "Times New Roman"
        style.font.name = face
        for slot in ("w:ascii", "w:hAnsi", "w:eastAsia"):
            style._element.rPr.rFonts.set(qn(slot), face)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)

    header = section.header.paragraphs[0]
    header.text = (
        "Western Digital (Nasdaq: WDC)  |  Q4 FY2026 Earnings Update"
        if lang == "en"
        else "西部数据（纳斯达克：WDC）｜2026 财年第四季度业绩更新"
    )
    set_run_font(header.runs[0], lang, size=7.6, bold=True, color=BLACK, heading=True)
    header.paragraph_format.space_after = Pt(0)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("WDC Research  |  15 September 2026  |  ")
    set_run_font(r, lang, size=7.5, color=GREY)
    page_field(footer)
    return doc


def new_page(doc):
    doc.add_page_break()


def title_page(doc, lang, market):
    price = market["price"]
    market_cap = f"${market['market_cap'] / 1e9:.1f}bn" if market["market_cap"] != "N/A" else "N/A"
    range_text = f"${market['low']:.2f}-${market['high']:.2f}" if market["low"] != "N/A" else "N/A"
    upside = (600 / price - 1) * 100 if isinstance(price, (int, float)) else 0
    if lang == "en":
        paragraph(doc, "WESTERN DIGITAL CORPORATION", lang, style="Title", align=WD_ALIGN_PARAGRAPH.CENTER, after=1.0, heading=True)
        paragraph(doc, "Q4 FY2026 EARNINGS UPDATE", lang, size=13.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=1.0, heading=True)
        paragraph(doc, "Pricing and capacity gains beat the quarter while the post earnings reset improves risk reward", lang, size=10.4, bold=True, color=ORANGE, align=WD_ALIGN_PARAGRAPH.CENTER, after=5.0, heading=True)
        meta = [
            "Published: 15 September 2026 | Results released: 5 August 2026 | Delayed reaction update: 41 days after release",
            "Rating: BUY | Price target: $600 | Initiating post-separation quarterly coverage",
            f"Reference price: ${price:.2f} | Implied upside: {upside:.1f}% | Market capitalization: {market_cap}",
            f"52-week range: {range_text} | Longbridge live cross-check: ${market['longbridge_last']:.2f}" if isinstance(market["longbridge_last"], (int, float)) else f"52-week range: {range_text}",
        ]
        headers = ["Metric", "Reported", "Pre-results consensus", "Variance"]
        rows = [
            ["Revenue", "$3.747bn", "$3.700bn", "+$47m / +1.3%"],
            ["Non-GAAP diluted EPS", "$3.56", "$3.31", "+$0.25 / +7.6%"],
            ["Non-GAAP gross margin", "54.4%", "51.5% prior guide midpoint", "+290 bps"],
            ["Q1 FY2027 revenue guide", "$4.1bn +/-$0.1bn", "$3.9bn", "+$0.2bn / +5.1%"],
        ]
    else:
        paragraph(doc, "西部数据公司", lang, style="Title", align=WD_ALIGN_PARAGRAPH.CENTER, after=1.0, heading=True)
        paragraph(doc, "2026 财年第四季度业绩更新报告", lang, size=13.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=1.0, heading=True)
        paragraph(doc, "定价与容量提升推动业绩超预期 财报后估值回落改善风险收益比", lang, size=10.4, bold=True, color=ORANGE, align=WD_ALIGN_PARAGRAPH.CENTER, after=5.0, heading=True)
        meta = [
            "发布日期：2026 年 9 月 15 日｜财报发布：2026 年 8 月 5 日｜延迟反应更新：发布后第 41 天",
            "评级：买入｜目标价：600 美元｜首次按拆分后纯 HDD 口径建立季度覆盖",
            f"参考股价：{price:.2f} 美元｜潜在上涨空间：{upside:.1f}%｜市值：{market_cap}",
            f"52 周区间：{range_text}｜Longbridge 实时交叉核验：{market['longbridge_last']:.2f} 美元" if isinstance(market["longbridge_last"], (int, float)) else f"52 周区间：{range_text}",
        ]
        headers = ["指标", "实际值", "财报前一致预期", "差异"]
        rows = [
            ["营业收入", "37.47 亿美元", "37.00 亿美元", "+0.47 亿 / +1.3%"],
            ["Non-GAAP 摊薄 EPS", "3.56 美元", "3.31 美元", "+0.25 / +7.6%"],
            ["Non-GAAP 毛利率", "54.4%", "此前指引中值 51.5%", "+290 个基点"],
            ["FY2027 Q1 收入指引", "41 亿 +/-1 亿美元", "39 亿美元", "+2 亿 / +5.1%"],
        ]
    for line in meta:
        paragraph(doc, line, lang, size=8.5, bold=("Rating" in line or "评级" in line), after=1.2)
    heading(doc, "EARNINGS SCORECARD" if lang == "en" else "业绩计分卡", lang, 2)
    make_table(doc, headers, rows, lang, [1.75, 1.55, 2.05, 1.55], variance_col=3, font_size=7.5)
    source(doc, [
        ("WDC Q4 FY2026 release, 5 Aug. 2026" if lang == "en" else "WDC FY2026 Q4 财报公告，2026 年 8 月 5 日", URLS["release"]),
        ("MarketBeat consensus snapshot, 5 Aug. 2026" if lang == "en" else "MarketBeat 财报时点一致预期，2026 年 8 月 5 日", URLS["consensus"]),
    ], lang)
    heading(doc, "INVESTMENT TAKEAWAYS" if lang == "en" else "投资要点", lang, 2)
    if lang == "en":
        bullet(doc, "The quarter beat on both reported earnings and prior guidance.", "Revenue exceeded consensus by $47m and the prior guide midpoint by $97m. Non-GAAP EPS beat consensus by $0.25 and the prior midpoint by $0.31. The 54.4% gross margin was the more important surprise because it cleared the prior 51%-52% range by 240-340 bps.", lang, compact=True)
        bullet(doc, "Pricing has become the main earnings lever.", "Management described blended price per terabyte growth in the high teens while cost per terabyte fell about 8%. Higher-capacity drives, LTA repricing and disciplined unit supply converted 44% revenue growth into 109% non-GAAP EPS growth.", lang, compact=True)
        bullet(doc, "The Q1 FY2027 guide supports another estimate reset.", "The $4.1bn revenue midpoint was about 5% above the contemporaneous Street figure and the $4.00 EPS midpoint was about 10% above consensus. We lift our FY2027 revenue and EPS framework to $19.6bn and $20.50.", lang, compact=True)
        bullet(doc, "Buy with a $600 target after the post-release drawdown.", "The shares fell about 13% on the first full session after results and trade below the release-day close. Our target blends an earnings multiple and a cash-flow valuation, while explicitly discounting customer concentration, technology-transition execution and LTA rigidity.", lang, compact=True)
    else:
        bullet(doc, "本季实际业绩及此前指引均被超越。", "收入较一致预期高 4,700 万美元，较此前指引中值高 9,700 万美元；Non-GAAP EPS 分别高 0.25 美元和 0.31 美元。54.4% 毛利率更为关键，较此前 51%-52% 区间高 240-340 个基点。", lang, compact=True)
        bullet(doc, "定价已成为盈利改善的核心杠杆。", "管理层称每 TB 综合售价同比增幅达到十几个百分点的高段，而每 TB 成本下降约 8%。高容量硬盘、LTA 重新定价与克制的单位供给，使 44% 收入增长转化为 109% Non-GAAP EPS 增长。", lang, compact=True)
        bullet(doc, "FY2027 Q1 指引支持再次上修预测。", "41 亿美元收入中值较当时一致预期高约 5%，4.00 美元 EPS 中值高约 10%。我们将 FY2027 收入和 EPS 预测上调至 196 亿美元与 20.50 美元。", lang, compact=True)
        bullet(doc, "财报后回撤带来买入机会，目标价 600 美元。", "财报后的首个完整交易日股价下跌约 13%，目前低于发布日收盘价。目标价综合盈利倍数与现金流估值，并对客户集中、技术切换执行以及 LTA 刚性给予明确折价。", lang, compact=True)


def page_results(doc, lang):
    new_page(doc)
    heading(doc, "RESULTS AND EARNINGS QUALITY" if lang == "en" else "业绩表现与盈利质量", lang, 1)
    if lang == "en":
        paragraph(doc, "Q4 revenue reached $3.747bn, up 12% sequentially and 44% year over year. The result was $47m above the $3.70bn pre-results consensus and near the top of management's $3.55bn-$3.75bn range. Growth remained concentrated in Cloud, but Client and Consumer also expanded as stronger pricing carried across the portfolio. The quarter completed a fiscal year in which revenue rose 36% to $12.919bn.", lang)
        paragraph(doc, "Non-GAAP diluted EPS rose to $3.56 from $2.72 in Q3 and $1.70 a year earlier. The $0.25 consensus beat was larger than the revenue surprise because gross margin expanded faster than expected and operating expenses declined 4% sequentially to $382m. Non-GAAP operating income reached $1.655bn, producing a 44.2% operating margin. The result shows that incremental revenue is currently converting at unusually high rates.", lang)
        paragraph(doc, "GAAP diluted EPS of $8.21 is not the right run-rate measure. It includes a $2.05bn mark-to-market gain on the retained Sandisk interest and $362m of costs linked to debt and equity transactions. We therefore anchor operating comparisons and valuation to non-GAAP EPS, while using GAAP cash flow and balance-sheet disclosures to test earnings quality.", lang)
    else:
        paragraph(doc, "Q4 收入达到 37.47 亿美元，环比增长 12%，同比增长 44%。该数字较 37.00 亿美元财报前一致预期高 4,700 万美元，也接近管理层 35.5 亿至 37.5 亿美元指引区间上沿。增长仍主要来自 Cloud，但 Client 与 Consumer 也受益于全产品组合定价改善。FY2026 全年收入因此增长 36% 至 129.19 亿美元。", lang)
        paragraph(doc, "Non-GAAP 摊薄 EPS 从 Q3 的 2.72 美元与上年同期的 1.70 美元升至 3.56 美元。相对一致预期的 0.25 美元超额幅度高于收入超额，因为毛利率扩张快于预期，同时营业费用环比下降 4% 至 3.82 亿美元。Non-GAAP 营业利润达到 16.55 亿美元，营业利润率为 44.2%，说明当前增量收入正以异常高的比例转化为利润。", lang)
        paragraph(doc, "8.21 美元 GAAP 摊薄 EPS 不能代表持续经营能力，其中包含所持 Sandisk 股权按市值计价产生的 20.50 亿美元收益，以及与债务和股权交易相关的 3.62 亿美元成本。因此，经营比较与估值以 Non-GAAP EPS 为核心，同时用 GAAP 现金流和资产负债表披露检验盈利质量。", lang)
    figure(doc, 1, "Quarterly revenue progression" if lang == "en" else "季度收入走势", lang, [
        ("Q4 FY2026 presentation, slides 5-6" if lang == "en" else "FY2026 Q4 演示材料，第 5-6 页", URLS["presentation"]),
        ("Q3 FY2026 presentation, slides 5-6" if lang == "en" else "FY2026 Q3 演示材料，第 5-6 页", URLS["q3_presentation"]),
    ], width=5.35)
    figure(doc, 2, "Non-GAAP diluted EPS progression" if lang == "en" else "Non-GAAP 摊薄 EPS 走势", lang, [
        ("Q4 FY2026 presentation, slides 6 and 11" if lang == "en" else "FY2026 Q4 演示材料，第 6 页与第 11 页", URLS["presentation"]),
        ("Q3 FY2026 presentation, slides 6 and 10" if lang == "en" else "FY2026 Q3 演示材料，第 6 页与第 10 页", URLS["q3_presentation"]),
    ], width=5.35)


def page_profitability(doc, lang):
    new_page(doc)
    heading(doc, "MARGINS CASH FLOW AND BALANCE SHEET" if lang == "en" else "利润率 现金流与资产负债表", lang, 1)
    if lang == "en":
        paragraph(doc, "Non-GAAP gross margin expanded 390 bps sequentially and 1,310 bps year over year to 54.4%. The earnings call identified three drivers: richer mix from high-capacity nearline products, favorable pricing across Cloud and non-nearline markets, and manufacturing execution. The result matters because it was well above the 51%-52% prior guide and because management guided another step to 55%-56% in Q1 FY2027.", lang)
        paragraph(doc, "Operating leverage amplified the gross-profit gain. Non-GAAP operating expenses were $382m, down from $397m in Q3 despite a 12% revenue increase. Non-GAAP operating margin consequently rose 560 bps sequentially to 44.2%. This margin profile should not be extrapolated without qualification: HDD pricing is cyclical, LTA resets occur at different times, and the 53-week FY2026 calendar gave Q1 an extra week.", lang)
        paragraph(doc, "Cash conversion was strong. Operating cash flow reached $1.389bn and capital expenditure was $108m, yielding $1.281bn of free cash flow and a 34% FCF margin. FY2026 FCF totaled $3.511bn, up 145%. WDC ended July with $1.579bn of cash and $1.052bn of debt, a net cash position of about $527m. Accounts receivable increased to $2.026bn and inventory to $1.511bn, so working-capital discipline remains relevant even with strong current cash generation.", lang)
    else:
        paragraph(doc, "Non-GAAP 毛利率环比提高 390 个基点、同比提高 1,310 个基点至 54.4%。电话会将改善归因于三项因素：高容量近线产品带来的组合优化、Cloud 与非近线市场的有利定价，以及制造执行。该结果重要之处在于显著高于此前 51%-52% 指引，管理层同时预计 FY2027 Q1 将进一步升至 55%-56%。", lang)
        paragraph(doc, "营业杠杆放大了毛利增长。Non-GAAP 营业费用为 3.82 亿美元，低于 Q3 的 3.97 亿美元，而收入环比增长 12%；营业利润率因而环比提高 560 个基点至 44.2%。但不宜无条件外推这一利润率：HDD 定价具有周期性，不同 LTA 的重定价时点各异，且 FY2026 为 53 周财年，Q1 多出一周。", lang)
        paragraph(doc, "现金转换表现强劲。经营现金流为 13.89 亿美元，资本开支 1.08 亿美元，自由现金流达到 12.81 亿美元，对应 34% FCF 利润率。FY2026 自由现金流为 35.11 亿美元，同比增长 145%。7 月末现金 15.79 亿美元、债务 10.52 亿美元，净现金约 5.27 亿美元；应收账款增至 20.26 亿美元、库存增至 15.11 亿美元，因此营运资金纪律仍需关注。", lang)
    figure(doc, 3, "Non-GAAP margin progression" if lang == "en" else "Non-GAAP 利润率走势", lang, [
        ("Q4 FY2026 presentation, slides 5-6 and 10" if lang == "en" else "FY2026 Q4 演示材料，第 5-6 页与第 10 页", URLS["presentation"]),
    ], width=5.35)
    figure(doc, 4, "Quarterly free cash flow progression" if lang == "en" else "季度自由现金流走势", lang, [
        ("Q4 FY2026 release, pages 8-10" if lang == "en" else "FY2026 Q4 财报公告，第 8-10 页", URLS["release"]),
        ("Q3 FY2026 presentation, slide 10" if lang == "en" else "FY2026 Q3 演示材料，第 10 页", URLS["q3_presentation"]),
    ], width=5.35)


def page_metrics(doc, lang):
    new_page(doc)
    heading(doc, "END MARKETS AND OPERATING METRICS" if lang == "en" else "终端市场与经营指标", lang, 1)
    if lang == "en":
        paragraph(doc, "Cloud revenue was about $3.3bn and represented 89% of Q4 sales. Client contributed $225m, or 6%, and Consumer contributed $187m, or 5%. All three end markets grew year over year, but the investment case is fundamentally a Cloud storage case: the 10-K shows Cloud generated $11.49bn of FY2026 revenue, 89% of the total, compared with $726m from Client and $703m from Consumer.", lang)
        paragraph(doc, "WDC shipped 231 exabytes in Q4, up 22% year over year. Nearline shipments increased to 209 EB from 170 EB, while non-nearline was 22 EB. The 22% total growth rate slowed from recent quarters because product and customer mix is lumpy, not because management identified a demand break. Management expects demand and shipments to grow above 25% over time as 40 TB ePMR ramps and 44 TB HAMR enters the market in calendar 2027.", lang)
        paragraph(doc, "Technology cadence is therefore central to the next phase. WDC began shipping next-generation ePMR drives at up to 40 TB in Q4 and expects them to exceed half of nearline exabytes by Q3 FY2027. UltraSMR is expected to reach about 60% of nearline exabytes by the end of FY2027. High-bandwidth drives were sampling with five customers and offered up to eight times the throughput of conventional designs, potentially extending HDD economics into more data-intensive AI workloads.", lang)
    else:
        paragraph(doc, "Cloud 收入约 33 亿美元，占 Q4 销售额 89%；Client 为 2.25 亿美元，占 6%；Consumer 为 1.87 亿美元，占 5%。三类终端市场均同比增长，但投资逻辑本质上仍是云存储逻辑：10-K 显示 FY2026 Cloud 收入 114.90 亿美元，占总收入 89%，Client 与 Consumer 分别仅为 7.26 亿和 7.03 亿美元。", lang)
        paragraph(doc, "Q4 出货量为 231 EB，同比增长 22%。其中近线产品从 170 EB 增至 209 EB，非近线产品为 22 EB。22% 总增速较近期季度放缓，管理层将其归因于产品和客户组合的季度波动，而不是需求断裂。随着 40 TB ePMR 放量、44 TB HAMR 于 2027 日历年进入市场，公司预计中长期需求与出货量增速可保持在 25% 以上。", lang)
        paragraph(doc, "因此，技术节奏决定下一阶段兑现程度。WDC 在 Q4 开始出货最高 40 TB 的新一代 ePMR，并预计到 FY2027 Q3 占近线 EB 的比例超过一半；UltraSMR 到 FY2027 末预计占近线 EB 约 60%。高带宽硬盘已向五家客户送样，吞吐量可达传统设计的八倍，有望把 HDD 的经济性延伸至数据强度更高的 AI 工作负载。", lang)
    figure(doc, 5, "Quarterly revenue mix by end market" if lang == "en" else "按终端市场划分的季度收入结构", lang, [
        ("Q4 FY2026 presentation, slide 4" if lang == "en" else "FY2026 Q4 演示材料，第 4 页", URLS["presentation"]),
        ("FY2026 Form 10-K, revenue by end market" if lang == "en" else "FY2026 Form 10-K，按终端市场划分的收入", URLS["tenk_doc"]),
    ], width=5.35)
    figure(doc, 6, "Nearline and non-nearline exabytes shipped" if lang == "en" else "近线与非近线 EB 出货量", lang, [
        ("Q4 FY2026 presentation, slide 5" if lang == "en" else "FY2026 Q4 演示材料，第 5 页", URLS["presentation"]),
        ("Q4 FY2026 earnings call, 5 Aug. 2026" if lang == "en" else "FY2026 Q4 电话会，2026 年 8 月 5 日", URLS["transcript_full"]),
    ], width=5.35)


def page_guidance(doc, lang):
    new_page(doc)
    heading(doc, "GUIDANCE AND NEAR TERM OUTLOOK" if lang == "en" else "指引与近期展望", lang, 1)
    if lang == "en":
        paragraph(doc, "Q4 delivered above the prior guidance midpoint on every major operating line. Revenue beat the $3.65bn midpoint by 2.7%, non-GAAP gross margin beat the 51.5% midpoint by 290 bps, and non-GAAP EPS beat the $3.25 midpoint by 9.5%. These variances were driven by pricing, high-capacity mix and expense leverage rather than a one-time revenue recognition item.", lang)
        headers = ["Metric", "Q4 prior guide", "Q4 actual", "Q1 FY2027 guide", "Street before Q1"]
        rows = [
            ["Revenue", "$3.65bn +/-$0.10bn", "$3.747bn", "$4.1bn +/-$0.1bn", "$3.9bn"],
            ["Non-GAAP gross margin", "51%-52%", "54.4%", "55%-56%", "Not disclosed"],
            ["Non-GAAP operating expenses", "$385m-$395m", "$382m", "$390m-$400m", "Not disclosed"],
            ["Non-GAAP diluted EPS", "$3.25 +/-$0.15", "$3.56", "$4.00 +/-$0.15", "$3.63"],
        ]
    else:
        paragraph(doc, "Q4 各项主要经营指标均高于此前指引中值。收入较 36.5 亿美元中值高 2.7%，Non-GAAP 毛利率较 51.5% 中值高 290 个基点，Non-GAAP EPS 较 3.25 美元中值高 9.5%。这些差异来自定价、高容量产品组合与费用杠杆，而不是一次性收入确认。", lang)
        headers = ["指标", "Q4 此前指引", "Q4 实际", "FY2027 Q1 指引", "Q1 财报前市场预期"]
        rows = [
            ["收入", "36.5 亿 +/-1 亿美元", "37.47 亿美元", "41 亿 +/-1 亿美元", "39 亿美元"],
            ["Non-GAAP 毛利率", "51%-52%", "54.4%", "55%-56%", "未披露"],
            ["Non-GAAP 营业费用", "3.85-3.95 亿美元", "3.82 亿美元", "3.90-4.00 亿美元", "未披露"],
            ["Non-GAAP 摊薄 EPS", "3.25 +/-0.15 美元", "3.56 美元", "4.00 +/-0.15 美元", "3.63 美元"],
        ]
    make_table(doc, headers, rows, lang, [1.55, 1.55, 1.15, 1.55, 1.35], variance_col=None, font_size=7.15)
    source(doc, [
        ("Q3 FY2026 presentation, slide 7" if lang == "en" else "FY2026 Q3 演示材料，第 7 页", URLS["q3_presentation"]),
        ("Q4 FY2026 release, page 3" if lang == "en" else "FY2026 Q4 财报公告，第 3 页", URLS["release"]),
        ("MarketBeat consensus snapshot" if lang == "en" else "MarketBeat 一致预期快照", URLS["consensus"]),
    ], lang)
    if lang == "en":
        paragraph(doc, "The Q1 FY2027 outlook calls for 9% sequential revenue growth at the midpoint, a further 110 bps of gross-margin expansion, and 12% sequential EPS growth. The midpoint exceeded the contemporaneous Street view by roughly $200m of revenue and $0.37 of EPS. The guide therefore implies that price realization and mix remain favorable even as operating expenses rise modestly to support the product roadmap.", lang)
        paragraph(doc, "We view the guide as achievable but not conservative. WDC is relying on multi-quarter demand visibility and a disciplined supply environment, while individual hyperscaler purchase patterns can shift between quarters. The key monitor is whether exabyte growth reaccelerates above 25% as 40 TB products ramp; a guide beat driven only by price would still support earnings, but it would be less durable than simultaneous volume and price gains.", lang)
    else:
        paragraph(doc, "FY2027 Q1 指引中值对应收入环比增长 9%、毛利率再提高 110 个基点、EPS 环比增长 12%。该中值较当时市场预期高约 2 亿美元收入与 0.37 美元 EPS，意味着即便为产品路线图适度增加费用，定价兑现与产品组合仍然有利。", lang)
        paragraph(doc, "我们认为指引可实现，但并不保守。WDC 依赖多季度需求可见度与克制的行业供给，而单个超大规模客户的采购节奏可能在季度间移动。核心观察指标是 40 TB 产品放量后 EB 增速能否重新超过 25%；如果超预期仅由定价推动，仍有利于盈利，但耐久性弱于量价齐升。", lang)
    figure(doc, 7, "Q4 actual results compared with prior guidance" if lang == "en" else "Q4 实际业绩与此前指引比较", lang, [
        ("Q3 FY2026 presentation, slide 7" if lang == "en" else "FY2026 Q3 演示材料，第 7 页", URLS["q3_presentation"]),
        ("Q4 FY2026 release, 5 Aug. 2026" if lang == "en" else "FY2026 Q4 财报公告，2026 年 8 月 5 日", URLS["release"]),
    ], width=5.7)


def page_thesis(doc, lang):
    new_page(doc)
    heading(doc, "UPDATED INVESTMENT THESIS" if lang == "en" else "更新后的投资逻辑", lang, 1)
    if lang == "en":
        bullet(doc, "Cloud data growth and AI persistence strengthened.", "Cloud remained 89% of revenue and grew about 43% year over year. Management is discussing LTAs that extend into calendar 2029-2031, suggesting customers are planning storage capacity well beyond the current AI infrastructure cycle. Training, inference, agentic AI, physical AI and conventional video workloads all create persistent data that is economically suited to HDD storage.", lang)
        bullet(doc, "Pricing discipline strengthened.", "Blended price per terabyte improved from high-single-digit year-over-year growth in Q3 to the high teens in Q4. Meanwhile, cost per terabyte declined about 8%. This positive spread explains much of the 1,310 bps gross-margin expansion and demonstrates that higher-capacity drives can create value for customers and suppliers simultaneously.", lang)
        bullet(doc, "Technology execution remains on track but is not complete.", "The 40 TB ePMR launch began in Q4, with a target of more than 50% of nearline exabytes by Q3 FY2027. The next proof points are 44 TB HAMR in the first half of calendar 2027 and 50 TB products later that year. Delays in qualification, yield or customer deployment would weaken both volume and margin assumptions.", lang)
        bullet(doc, "Cash return and balance-sheet repair strengthened.", "FY2026 FCF reached $3.511bn and the company ended with net cash. Management said it remains committed to returning free cash flow through dividends and repurchases. The stronger balance sheet gives WDC room to fund heads, media and automation while continuing capital returns, but repurchases should be judged against the prevailing valuation rather than treated as automatically accretive.", lang)
        paragraph(doc, "The thesis is now less about an ordinary cyclical HDD recovery and more about whether WDC can sustain structurally better pricing under long-duration customer commitments. The difference is important. A normal recovery would justify a mid-cycle multiple on normalized earnings; durable LTA-backed demand, rising areal density and high incremental margins justify a higher forward multiple. We use the second framework but retain a material discount for concentration and contract rigidity.", lang)
    else:
        bullet(doc, "云数据增长与 AI 持久化需求得到强化。", "Cloud 继续占收入 89%，同比增长约 43%。管理层正讨论延伸至 2029-2031 日历年的 LTA，说明客户规划的存储容量远超当前 AI 基础设施周期。训练、推理、智能体 AI、物理 AI 与传统视频工作负载都会产生需要长期保存、且适合用 HDD 经济存储的数据。", lang)
        bullet(doc, "定价纪律得到强化。", "每 TB 综合售价同比增速从 Q3 的高个位数提升至 Q4 的十几个百分点高段，而每 TB 成本下降约 8%。这一正向价差解释了毛利率同比提高 1,310 个基点的大部分原因，也表明高容量硬盘能够同时为客户与供应商创造价值。", lang)
        bullet(doc, "技术执行保持正轨，但尚未完成。", "40 TB ePMR 已于 Q4 开始出货，目标是在 FY2027 Q3 前占近线 EB 超过 50%。下一验证点是 44 TB HAMR 于 2027 日历年上半年推出，以及 50 TB 产品在当年晚些时候推出。认证、良率或客户部署延迟都会削弱销量与利润率假设。", lang)
        bullet(doc, "现金回报与资产负债表修复得到强化。", "FY2026 自由现金流达到 35.11 亿美元，公司年末转为净现金。管理层仍承诺通过股息和回购返还自由现金流。更强的资产负债表能够同时支持磁头、介质和自动化投资及资本回报，但回购是否增值仍取决于当时估值。", lang)
        paragraph(doc, "投资逻辑已从普通 HDD 周期复苏，转向 WDC 能否在长期客户承诺下维持结构性更好的定价。两者估值含义不同：普通复苏只应给予正常化盈利的周期中段倍数；由 LTA 支撑的耐久需求、面密度提升与高增量利润率，则可以支持更高的前瞻倍数。我们采用第二种框架，但仍因集中度与合同刚性给予明显折价。", lang)
    figure(doc, 8, "Year-over-year price and cost movement per terabyte" if lang == "en" else "每 TB 售价与成本的同比变化", lang, [
        ("Q4 FY2026 earnings call, management remarks and Q&A" if lang == "en" else "FY2026 Q4 电话会，管理层陈述与问答", URLS["transcript_full"]),
    ], width=5.8)


def page_estimates(doc, lang, market):
    new_page(doc)
    heading(doc, "UPDATED ESTIMATES" if lang == "en" else "更新后的预测", lang, 1)
    price = float(market["price"]) if isinstance(market["price"], (int, float)) else 415.82
    old_pe = price / 17.50
    new_pe = price / 20.50
    fy28_pe = price / 32.50
    if lang == "en":
        paragraph(doc, "WDC has no prior company-specific report in this repository, so the 'old' column is a reconstructed pre-Q4 baseline anchored to the Q3 guidance trajectory and the pre-results Street snapshot. It is not a previously published house estimate. The 'new' column is our independent post-results view and is not company guidance beyond Q1 FY2027.", lang)
        headers = ["Metric", "FY2027E old", "FY2027E new", "Change", "FY2028E new"]
        rows = [
            ["Revenue ($bn)", "18.1", "19.6", "+8.3%", "26.8"],
            ["Revenue growth", "40.1%", "51.7%", "+1,160 bps", "36.7%"],
            ["Non-GAAP gross margin", "52.0%", "55.5%", "+350 bps", "56.5%"],
            ["Non-GAAP operating income ($bn)", "7.2", "8.6", "+19.4%", "12.2"],
            ["Non-GAAP operating margin", "39.8%", "43.9%", "+410 bps", "45.5%"],
            ["Non-GAAP diluted EPS", "$17.50", "$20.50", "+17.1%", "$32.50"],
            ["Free cash flow ($bn)", "4.6", "5.4", "+17.4%", "8.0"],
            [f"P/E at ${price:.2f}", f"{old_pe:.1f}x", f"{new_pe:.1f}x", f"{new_pe-old_pe:.1f}x", f"{fy28_pe:.1f}x"],
        ]
    else:
        paragraph(doc, "本仓库此前没有 WDC 专项报告，因此“旧预测”是以 Q3 指引路径和财报前市场快照重建的 Q4 前基准，并非曾经发布的内部预测；“新预测”是我们根据本次业绩形成的独立判断，除 FY2027 Q1 外均不属于公司指引。", lang)
        headers = ["指标", "FY2027E 旧", "FY2027E 新", "变动", "FY2028E 新"]
        rows = [
            ["收入（十亿美元）", "18.1", "19.6", "+8.3%", "26.8"],
            ["收入增速", "40.1%", "51.7%", "+1,160 个基点", "36.7%"],
            ["Non-GAAP 毛利率", "52.0%", "55.5%", "+350 个基点", "56.5%"],
            ["Non-GAAP 营业利润（十亿美元）", "7.2", "8.6", "+19.4%", "12.2"],
            ["Non-GAAP 营业利润率", "39.8%", "43.9%", "+410 个基点", "45.5%"],
            ["Non-GAAP 摊薄 EPS", "17.50 美元", "20.50 美元", "+17.1%", "32.50 美元"],
            ["自由现金流（十亿美元）", "4.6", "5.4", "+17.4%", "8.0"],
            [f"按 {price:.2f} 美元计算 P/E", f"{old_pe:.1f} 倍", f"{new_pe:.1f} 倍", f"{new_pe-old_pe:.1f} 倍", f"{fy28_pe:.1f} 倍"],
        ]
    make_table(doc, headers, rows, lang, [2.25, 1.25, 1.25, 1.2, 1.25], variance_col=3, font_size=7.15)
    source(doc, [
        ("Our estimates" if lang == "en" else "本报告预测", None),
        ("Q4 FY2026 release and Q1 guide" if lang == "en" else "FY2026 Q4 财报公告与 Q1 指引", URLS["release"]),
        ("Yahoo Finance consensus as of 15 Sep. 2026" if lang == "en" else "Yahoo Finance 一致预期，截至 2026 年 9 月 15 日", URLS["yahoo"]),
    ], lang)
    if lang == "en":
        paragraph(doc, "Our FY2027 revenue is slightly above the current Yahoo Finance consensus near $19.2bn. We model Q1 close to the $4.1bn guide midpoint, followed by stronger exabyte growth as 40 TB products ramp and LTA pricing resets. The 55.5% full-year gross margin assumes the Q1 midpoint is broadly sustained, with further product benefits offset by normal quarterly mix variation.", lang)
        paragraph(doc, "FY2028 assumes 36.7% revenue growth, a 100 bps gross-margin gain and operating expenses growing well below revenue. This produces $32.50 of non-GAAP EPS and $8.0bn of FCF. The assumptions are deliberately below a simple continuation of the current quarterly revenue growth rate, but they still require strong cloud capital spending, orderly industry supply and timely HAMR execution.", lang)
    else:
        paragraph(doc, "FY2027 收入预测略高于 Yahoo Finance 当前约 191.9 亿美元的一致预期。我们假设 Q1 接近 41 亿美元指引中值，随后 40 TB 产品放量与 LTA 重定价推动 EB 增速提高。全年 55.5% 毛利率假设 Q1 中值大致可维持，后续产品收益被正常的季度组合波动部分抵消。", lang)
        paragraph(doc, "FY2028 假设收入增长 36.7%、毛利率提高 100 个基点，营业费用增速显著低于收入，从而实现 32.50 美元 Non-GAAP EPS 与 80 亿美元自由现金流。这些假设低于简单延续当前季度收入增速的结果，但仍需要云资本开支强劲、行业供给有序、HAMR 按期执行。", lang)
    figure(doc, 9, "FY2027 estimate revisions" if lang == "en" else "FY2027 预测上修", lang, [
        ("Our estimates and reconstructed pre-Q4 baseline" if lang == "en" else "本报告预测与重建的 Q4 前基准", None),
        ("Q4 FY2026 earnings materials" if lang == "en" else "FY2026 Q4 财报材料", URLS["release"]),
    ], width=5.75)


def page_valuation(doc, lang, market):
    new_page(doc)
    heading(doc, "VALUATION AND PRICE TARGET" if lang == "en" else "估值与目标价", lang, 1)
    price = float(market["price"]) if isinstance(market["price"], (int, float)) else 415.82
    upside = (600 / price - 1) * 100
    if lang == "en":
        paragraph(doc, f"We initiate a Buy rating with a $600 price target, implying {upside:.1f}% upside from the ${price:.2f} reference price. The target is a rounded blend of two methods. A forward earnings approach applies 18.5x to FY2028E non-GAAP EPS of $32.50, producing $601. A DCF using FY2027-FY2031 FCF of $5.4bn, $8.0bn, $10.0bn, $12.5bn and $15.0bn, a 9.5% WACC, 3.5% terminal growth and current net cash produces about $561 per share. A 60% earnings and 40% DCF blend yields about $585; we round to $600 because current consensus target data and the company's accelerating earnings profile support the upper half of that range.", lang)
        paragraph(doc, "The multiple is below the 20x-23x range that the market has often assigned during the last year and below the current Longbridge consensus target of about $665. We do not use the trailing GAAP P/E because FY2026 GAAP income includes large Sandisk mark-to-market gains. On our estimates, the shares trade at about 20.3x FY2027 EPS and 12.8x FY2028 EPS. The apparent compression is attractive if the product and pricing assumptions hold, but it also signals how much earnings growth is already required.", lang)
        heading(doc, "Scenario framework", lang, 2)
        paragraph(doc, "The $360 bear case applies 12x to roughly $30 of FY2028 EPS. It assumes weaker cloud spending, slower HAMR qualification and price normalization. The $600 base applies 18.5x to $32.50. The $820 bull case applies 22x to about $37.30, requiring faster exabyte growth, sustained high-teens price realization and gross margin above 57%. The bear case is below the current price, so the rating depends on the probability of a structural rather than purely cyclical improvement.", lang)
    else:
        paragraph(doc, f"我们首次给予“买入”评级与 600 美元目标价，较 {price:.2f} 美元参考价有 {upside:.1f}% 上涨空间。目标价是两种方法的取整组合：前瞻盈利法对 FY2028E Non-GAAP EPS 32.50 美元采用 18.5 倍，得到 601 美元；DCF 假设 FY2027-FY2031 自由现金流依次为 54 亿、80 亿、100 亿、125 亿和 150 亿美元，WACC 为 9.5%、永续增长率 3.5%，并计入当前净现金，得到约 561 美元。60% 盈利法与 40% DCF 的组合约为 585 美元；考虑当前一致目标价与加速的盈利曲线位于区间上半部，我们取整至 600 美元。", lang)
        paragraph(doc, "18.5 倍低于过去一年市场常用的约 20-23 倍区间，也低于 Longbridge 当前约 665 美元的一致目标价。我们不使用过去十二个月 GAAP P/E，因为 FY2026 GAAP 利润包含大量 Sandisk 股权按市值计价收益。按本报告预测，当前股价对应 FY2027 与 FY2028 P/E 分别约 20.3 倍和 12.8 倍；若产品与定价假设兑现，这种快速压缩具有吸引力，但也说明市场仍要求显著盈利增长。", lang)
        heading(doc, "情景框架", lang, 2)
        paragraph(doc, "悲观情景 360 美元，对 FY2028 约 30 美元 EPS 采用 12 倍，假设云资本开支减弱、HAMR 认证放缓、定价正常化。基准情景 600 美元，对 32.50 美元采用 18.5 倍。乐观情景 820 美元，对约 37.30 美元采用 22 倍，需要 EB 增长更快、每 TB 售价持续实现十几个百分点高段增长、毛利率超过 57%。悲观价值低于当前股价，因此评级取决于盈利改善更可能是结构性而非纯周期性。", lang)
    figure(doc, 10, "Bear base and bull valuation scenarios" if lang == "en" else "悲观 基准与乐观估值情景", lang, [
        ("Our FY2028 estimates and valuation assumptions" if lang == "en" else "本报告 FY2028 预测与估值假设", None),
        ("Longbridge analyst target snapshot, 4 Sep. 2026" if lang == "en" else "Longbridge 分析师目标价快照，2026 年 9 月 4 日", URLS["longbridge"]),
        ("Yahoo Finance market data, 15 Sep. 2026" if lang == "en" else "Yahoo Finance 市场数据，2026 年 9 月 15 日", URLS["yahoo"]),
    ], width=5.85)


def page_risks(doc, lang, market):
    new_page(doc)
    heading(doc, "CATALYSTS RISKS AND TRADING CONTEXT" if lang == "en" else "催化剂 风险与交易背景", lang, 1)
    if lang == "en":
        heading(doc, "Catalysts", lang, 2)
        bullet(doc, "Q1 FY2027 execution in late October.", "Revenue above $4.1bn, gross margin near the upper half of 55%-56% and EPS above $4.00 would validate continued price realization. More important, management must show exabyte growth can accelerate as the 40 TB mix rises.", lang, compact=True)
        bullet(doc, "40 TB ePMR and 44 TB HAMR milestones.", "Crossing 50% of nearline exabytes with the next ePMR platform by Q3 FY2027 would improve capacity and cost per terabyte. HAMR qualification and initial deployment in calendar 2027 would extend the roadmap and reduce the risk that a competitor captures the technology premium.", lang, compact=True)
        bullet(doc, "Longer customer agreements and capital returns.", "Additional LTAs for 2029-2031 could improve demand visibility if pricing formulas preserve economics. Continued repurchases and dividends can support per-share value, provided buybacks do not outrun intrinsic value.", lang, compact=True)
        heading(doc, "Principal risks", lang, 2)
        bullet(doc, "Customer concentration.", "The top ten customers represented 73% of FY2026 revenue. Three customers contributed 16%, 15% and 13%; two represented 25% and 17% of receivables. A single purchasing change can therefore affect revenue, mix, working capital and factory utilization.", lang, RED, compact=True)
        bullet(doc, "LTA rigidity and HDD cyclicality.", "Long-term agreements improve visibility but can cap participation when market prices rise, and failures by either party can create volume, pricing or damage claims. Excess customer inventory or a cloud spending pause could still produce a sharp cyclical correction.", lang, RED, compact=True)
        bullet(doc, "Technology and supply execution.", "The estimate path assumes timely 40 TB ePMR, UltraSMR and HAMR ramps. Qualification delays, yield problems, component shortages or a rival's faster transition would reduce exabytes and cost improvement. WDC also depends on a limited supplier base.", lang, RED, compact=True)
        paragraph(doc, "The stock fell from an adjusted $519.01 close on the release date to $451.39 on 6 August, a 13.0% one-day decline, and closed near $416 on 15 September. The selloff occurred despite an earnings beat and strong guide, indicating expectations and relative margin comparisons had become demanding. We view the reset as useful, but the wide 52-week range confirms that position sizing matters.", lang)
    else:
        heading(doc, "催化剂", lang, 2)
        bullet(doc, "10 月下旬公布 FY2027 Q1。", "若收入高于 41 亿美元、毛利率位于 55%-56% 区间上半部、EPS 高于 4.00 美元，将验证定价持续兑现。更重要的是，公司需要证明 40 TB 产品占比提高后 EB 增速能够重新加速。", lang, compact=True)
        bullet(doc, "40 TB ePMR 与 44 TB HAMR 里程碑。", "到 FY2027 Q3，新一代 ePMR 占近线 EB 超过 50%，将同时改善容量与每 TB 成本。HAMR 在 2027 日历年的认证和初始部署将延伸路线图，并降低竞争对手独占技术溢价的风险。", lang, compact=True)
        bullet(doc, "更长期客户协议与资本回报。", "若价格公式能够保护经济性，新增 2029-2031 年 LTA 将提高需求可见度。持续回购和股息也可提升每股价值，前提是回购价格不高于内在价值。", lang, compact=True)
        heading(doc, "主要风险", lang, 2)
        bullet(doc, "客户集中。", "前十大客户占 FY2026 收入 73%，其中三家分别占 16%、15% 与 13%；两家客户分别占应收账款 25% 与 17%。单个客户采购变化可同时影响收入、组合、营运资金与工厂利用率。", lang, RED, compact=True)
        bullet(doc, "LTA 刚性与 HDD 周期性。", "长期协议提高可见度，但在市场价格上涨时可能限制公司参与上行；任何一方违约还可能导致销量、价格或赔偿问题。客户库存过高或云资本开支暂停仍可能带来剧烈周期回调。", lang, RED, compact=True)
        bullet(doc, "技术与供应执行。", "预测路径假设 40 TB ePMR、UltraSMR 与 HAMR 按期放量。认证延迟、良率问题、零部件短缺或竞争对手更快切换都会减少 EB 出货与成本改善，且 WDC 仍依赖有限的供应商群体。", lang, RED, compact=True)
        paragraph(doc, "经前复权后，股价从财报发布日收盘 519.01 美元跌至 8 月 6 日的 451.39 美元，单日下跌 13.0%，9 月 15 日收盘约 416 美元。业绩超预期且指引强劲仍遭抛售，说明市场预期与相对利润率比较此前已较苛刻。我们认为估值回落改善了机会，但宽阔的 52 周区间也意味着仓位管理很重要。", lang)
    figure(doc, 11, "Adjusted one-year share price and target" if lang == "en" else "前复权一年股价与目标价", lang, [
        ("Longbridge adjusted daily K-line through 15 Sep. 2026" if lang == "en" else "Longbridge 前复权日 K，截至 2026 年 9 月 15 日", URLS["longbridge"]),
        ("Yahoo Finance dynamic market data" if lang == "en" else "Yahoo Finance 动态市场数据", URLS["yahoo"]),
        ("Our valuation" if lang == "en" else "本报告估值", None),
    ], width=5.85)


def page_sources(doc, lang):
    new_page(doc)
    heading(doc, "SOURCES METHODOLOGY AND DISCLOSURES" if lang == "en" else "资料来源 方法与披露", lang, 1)
    if lang == "en":
        heading(doc, "Timeliness and issuer verification", lang, 2)
        paragraph(doc, "Today is 15 September 2026. Western Digital Corporation trades on Nasdaq under WDC and files with SEC CIK 106040. The company released Q4 and FY2026 results on 5 August 2026 for the period ended 3 July 2026; it furnished Form 8-K Item 2.02 the same day and filed Form 10-K on 14 August. The earnings call was held on 5 August at 4:30 p.m. Eastern, matching the release date. These are the latest reported results and were released 41 days before this delayed reaction update.", lang)
        heading(doc, "Comparability and source hierarchy", lang, 2)
        paragraph(doc, "Western Digital completed the separation of Sandisk on 21 February 2025. Historical figures used here follow WDC's continuing-operations presentation for the HDD business; they do not combine post-separation Sandisk flash revenue. Reported results, balance-sheet data, customer concentration and risk factors are anchored to the earnings release, presentation, Form 8-K and Form 10-K. The company webcast page and dated transcripts verify call timing and management commentary. Consensus uses MarketBeat's 5 August snapshot; current annual consensus and market fields use Yahoo Finance as of 15 September. Longbridge supplied the live quote cross-check, rating snapshot and adjusted K-line history.", lang)
        heading(doc, "Forecast conventions", lang, 2)
        paragraph(doc, "The pre-Q4 baseline is reconstructed because no earlier WDC coverage report exists in this repository. It uses the Q3 guidance path and pre-results consensus and should not be read as a published prior estimate. FY2027 and FY2028 figures are independent analytical estimates. Non-GAAP estimates follow WDC's continuing-operations definitions. The DCF uses nominal U.S. dollars and assumes the stated free-cash-flow path, 9.5% WACC, 3.5% terminal growth, current net cash and approximately 361m shares. Scenario values are judgments, not probability-weighted forecasts.", lang)
        heading(doc, "Important disclosure", lang, 2)
        paragraph(doc, "This report is independent analytical research for informational purposes only. It is not investment advice, an offer, a solicitation or a fiduciary recommendation. Rating, estimates, scenarios and price target are opinions based on assumptions that may prove wrong. Western Digital is exposed to customer and supplier concentration, HDD demand cycles, pricing changes, product qualification, manufacturing yield, technology transitions, long-term contract obligations, trade restrictions, cybersecurity, capital allocation and market valuation. Readers should review the primary filings and obtain advice suited to their circumstances.", lang)
        refs = [
            ("Q4 FY2026 earnings release, 5 Aug. 2026", URLS["release"]),
            ("Q4 FY2026 investor presentation, 5 Aug. 2026", URLS["presentation"]),
            ("Official Q4 FY2026 earnings webcast, 5 Aug. 2026", URLS["webcast"]),
            ("Q4 FY2026 earnings-call transcript, 5 Aug. 2026", URLS["transcript"]),
            ("Detailed call transcript and Q&A, 5 Aug. 2026", URLS["transcript_full"]),
            ("Form 8-K Item 2.02, filed 5 Aug. 2026", URLS["eightk"]),
            ("Form 10-K, filed 14 Aug. 2026", URLS["tenk"]),
            ("Q3 FY2026 release and prior guidance, 30 Apr. 2026", URLS["q3_release"]),
            ("Q3 FY2026 investor presentation, 30 Apr. 2026", URLS["q3_presentation"]),
            ("MarketBeat pre-results consensus snapshot", URLS["consensus"]),
            ("Yahoo Finance market data and consensus", URLS["yahoo"]),
            ("Longbridge WDC market data", URLS["longbridge"]),
            ("Western Digital Investor Relations", URLS["ir"]),
        ]
    else:
        heading(doc, "时效性与发行人核验", lang, 2)
        paragraph(doc, "今天是 2026 年 9 月 15 日。西部数据公司在纳斯达克以 WDC 交易，SEC CIK 为 106040。公司于 2026 年 8 月 5 日发布截至 7 月 3 日的 FY2026 Q4 与全年业绩，并在同日提交 Form 8-K Item 2.02，8 月 14 日提交 Form 10-K。电话会于 8 月 5 日美东时间 16:30 举行，与财报发布日期一致。这是最新已报告业绩，距本延迟反应更新发布 41 天。", lang)
        heading(doc, "可比口径与资料层级", lang, 2)
        paragraph(doc, "西部数据于 2025 年 2 月 21 日完成 Sandisk 分拆。本报告历史数据遵循 WDC 对 HDD 持续经营业务的重述口径，不合并分拆后 Sandisk 闪存收入。已报告业绩、资产负债表、客户集中与风险因素以财报公告、演示材料、Form 8-K 与 Form 10-K 为核心证据；公司 webcast 页面和带日期的文字实录用于核对电话会时间与管理层评论。一致预期采用 MarketBeat 8 月 5 日快照；当前年度一致预期和市场字段采用 9 月 15 日 Yahoo Finance；Longbridge 提供实时股价交叉核验、评级快照与前复权日 K。", lang)
        heading(doc, "预测口径", lang, 2)
        paragraph(doc, "由于本仓库此前没有 WDC 覆盖报告，Q4 前基准是依据 Q3 指引路径与财报前一致预期重建的，不应视为已经发布的旧预测。FY2027 与 FY2028 数字均为独立分析预测，Non-GAAP 预测遵循 WDC 持续经营业务定义。DCF 使用名义美元，假设正文所述自由现金流路径、9.5% WACC、3.5% 永续增长率、当前净现金与约 3.61 亿股。情景价值属于分析判断，而不是概率加权预测。", lang)
        heading(doc, "重要披露", lang, 2)
        paragraph(doc, "本报告为独立分析研究，仅供信息参考，不构成投资建议、要约、招揽或受托推荐。评级、预测、情景与目标价均基于可能不成立的假设。西部数据面临客户与供应商集中、HDD 需求周期、定价变化、产品认证、制造良率、技术切换、长期合同义务、贸易限制、网络安全、资本配置与市场估值等风险。读者应审阅原始申报文件，并取得适合自身情况的专业意见。", lang)
        refs = [
            ("FY2026 Q4 财报公告，2026 年 8 月 5 日", URLS["release"]),
            ("FY2026 Q4 投资者演示，2026 年 8 月 5 日", URLS["presentation"]),
            ("FY2026 Q4 官方电话会 webcast，2026 年 8 月 5 日", URLS["webcast"]),
            ("FY2026 Q4 电话会实录，2026 年 8 月 5 日", URLS["transcript"]),
            ("电话会详细陈述与问答，2026 年 8 月 5 日", URLS["transcript_full"]),
            ("Form 8-K Item 2.02，2026 年 8 月 5 日提交", URLS["eightk"]),
            ("Form 10-K，2026 年 8 月 14 日提交", URLS["tenk"]),
            ("FY2026 Q3 公告与此前指引，2026 年 4 月 30 日", URLS["q3_release"]),
            ("FY2026 Q3 投资者演示，2026 年 4 月 30 日", URLS["q3_presentation"]),
            ("MarketBeat 财报前一致预期快照", URLS["consensus"]),
            ("Yahoo Finance 市场数据与一致预期", URLS["yahoo"]),
            ("Longbridge WDC 市场数据", URLS["longbridge"]),
            ("西部数据投资者关系网站", URLS["ir"]),
        ]
    heading(doc, "Sources and references" if lang == "en" else "资料来源与参考文件", lang, 2)
    for label, url in refs:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2.1)
        r = p.add_run("• ")
        set_run_font(r, lang, size=8.2)
        add_hyperlink(p, label, url, lang, size=8.2)


def build(lang, market):
    doc = setup_doc(lang)
    title_page(doc, lang, market)
    page_results(doc, lang)
    page_profitability(doc, lang)
    page_metrics(doc, lang)
    page_guidance(doc, lang)
    page_thesis(doc, lang)
    page_estimates(doc, lang, market)
    page_valuation(doc, lang, market)
    page_risks(doc, lang, market)
    page_sources(doc, lang)
    filename = (
        "WDC_Q4_FY2026_Earnings_Update.docx"
        if lang == "en"
        else "WDC_Q4_FY2026_业绩更新报告_中文版.docx"
    )
    path = OUT / filename
    doc.save(path)
    return path


if __name__ == "__main__":
    missing = [str(path) for path in CHARTS.values() if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing chart files: " + ", ".join(missing))
    market = dynamic_market_data()
    print("Market data:", json.dumps(market, ensure_ascii=False))
    en_path = build("en", market)
    cn_path = build("cn", market)
    print("Created", en_path)
    print("Created", cn_path)
