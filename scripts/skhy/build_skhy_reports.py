#!/usr/bin/env python3
"""Build bilingual SK hynix Q2 2026 earnings-update reports."""

from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "output" / "SKHY"
OUT.mkdir(parents=True, exist_ok=True)

NAVY = "17365D"
BLUE = "2F5597"
PALE_BLUE = "DCE6F1"
ORANGE = "E67E22"
GREEN = "2E8B57"
RED = "C0392B"
GREY = "666666"
LIGHT = "F2F4F7"
WHITE = "FFFFFF"
BLACK = "111111"

URLS = {
    "release": "https://news.skhynix.com/en/q2-2026-business-results/",
    "presentation": "https://mis-prod-koce-homepage-cdn-01-blob-ep.azureedge.net/web/attach/13571040283022005.pdf",
    "call": "https://irsvc.teletogether.com/hynix/hynix2026Q2_eng.php",
    "ir": "https://www.skhynix.com/ir/UI-FR-IR01/",
    "sec6k": "https://www.sec.gov/Archives/edgar/data/2120882/000119312526303983/d19380d6k.htm",
    "prospectus": "https://www.sec.gov/Archives/edgar/data/2120882/000119312526299963/d32785d424b4.htm",
    "consensus": "https://en.yna.co.kr/view/AEN20260726000800320",
    "hanwha": "https://stock.mk.co.kr/uploads/20260622/1782105203_d85b7cf4d33e7d029f34.pdf",
    "q1": "https://mis-prod-koce-homepage-cdn-01-blob-ep.azureedge.net/web/attach/18333727317328069.pdf",
    "q425": "https://mis-prod-koce-homepage-cdn-01-blob-ep.azureedge.net/web/attach/116964051831620258.pdf",
    "q325": "https://mis-prod-koce-homepage-cdn-01-blob-ep.azureedge.net/web/attach/3127648149034921.pdf",
    "q225": "https://mis-prod-koce-homepage-cdn-01-blob-ep.azureedge.net/web/attach/40103454849895488.pdf",
    "q125": "https://mis-prod-koce-homepage-cdn-01-blob-ep.azureedge.net/web/attach/5001632852769028.pdf",
    "q424": "https://mis-prod-koce-homepage-cdn-01-blob-ep.azureedge.net/web/attach/81727201766385339.pdf",
    "q324": "https://mis-prod-koce-homepage-cdn-01-blob-ep.azureedge.net/web/attach/77045741902433810.pdf",
}

CHARTS = {i: OUT / f"skhy_chart{i}_{name}.png" for i, name in {
    1: "revenue", 2: "operating_profit", 3: "margins", 4: "normalized_net",
    5: "beat_miss", 6: "product_mix", 7: "cash_debt", 8: "cash_flow",
    9: "outlook", 10: "estimates", 11: "price", 12: "valuation",
}.items()}


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=110, bottom=80, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, twips):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(twips))
    tc_w.set(qn("w:type"), "dxa")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_fixed(table, widths):
    usable = 9994
    total = sum(widths)
    widths = [round(w * usable / total) for w in widths]
    widths[-1] += usable - sum(widths)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(usable))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "110")
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            set_cell_width(cell, widths[i])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def font_run(run, lang, size=9.6, bold=False, color=BLACK, italic=False, heading=False):
    latin = "Times New Roman"
    east = ("Heiti SC" if heading else "Songti SC") if lang == "cn" else "Times New Roman"
    run.font.name = east if lang == "cn" else latin
    run._element.rPr.rFonts.set(qn("w:ascii"), east if lang == "cn" else latin)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), east if lang == "cn" else latin)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    if lang == "cn":
        lang_el = run._element.rPr.find(qn("w:lang"))
        if lang_el is None:
            lang_el = OxmlElement("w:lang")
            run._element.rPr.append(lang_el)
        lang_el.set(qn("w:val"), "zh-CN")
        lang_el.set(qn("w:eastAsia"), "zh-CN")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_hyperlink(paragraph, text, url, lang, size=8):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_fonts = OxmlElement("w:rFonts")
    link_font = "Songti SC" if lang == "cn" else "Times New Roman"
    r_fonts.set(qn("w:ascii"), link_font)
    r_fonts.set(qn("w:hAnsi"), link_font)
    r_fonts.set(qn("w:eastAsia"), link_font)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    for item in (color, underline, r_fonts, sz):
        r_pr.append(item)
    if lang == "cn":
        lang_el = OxmlElement("w:lang")
        lang_el.set(qn("w:val"), "zh-CN")
        lang_el.set(qn("w:eastAsia"), "zh-CN")
        r_pr.append(lang_el)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    r._r.extend([fld_char1, instr, fld_char2])
    font_run(r, "en", 8, color=GREY)


def setup_doc(lang):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.70)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.78)
    sec.right_margin = Inches(0.78)
    sec.header_distance = Inches(0.30)
    sec.footer_distance = Inches(0.32)

    normal = doc.styles["Normal"]
    normal.font.name = "Songti SC" if lang == "cn" else "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Songti SC" if lang == "cn" else "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Songti SC" if lang == "cn" else "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Songti SC" if lang == "cn" else "Times New Roman")
    normal.font.size = Pt(9.6)
    normal.paragraph_format.space_after = Pt(4.5)
    normal.paragraph_format.line_spacing = 1.05

    for style_name, size, color in (("Title", 19, NAVY), ("Heading 1", 14, NAVY), ("Heading 2", 11.5, BLUE)):
        st = doc.styles[style_name]
        st.font.name = "Heiti SC" if lang == "cn" else "Times New Roman"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Heiti SC" if lang == "cn" else "Times New Roman")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Heiti SC" if lang == "cn" else "Times New Roman")
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Heiti SC" if lang == "cn" else "Times New Roman")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.keep_with_next = True

    header = sec.header.paragraphs[0]
    header.text = "SK hynix (Nasdaq: SKHY)  |  Q2 FY2026 Earnings Update" if lang == "en" else "SK hynix（纳斯达克：SKHY）｜2026 财年第二季度业绩更新"
    font_run(header.runs[0], lang, 8, bold=True, color=NAVY)
    header.paragraph_format.space_after = Pt(0)
    footer = sec.footer.paragraphs[0]
    footer.add_run("SKHY Research  |  2 August 2026     ")
    font_run(footer.runs[0], lang, 8, color=GREY)
    add_page_field(footer)
    return doc


def para(doc, text, lang, bold=False, size=9.6, color=BLACK, before=0, after=4.5, align=None, heading_font=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    font_run(r, lang, size, bold=bold, color=color, heading=heading_font)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.05
    if align is not None:
        p.alignment = align
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(2 if level == 1 else 1)
    p.paragraph_format.space_after = Pt(5)
    return p


def bullet(doc, title, body, lang, color=NAVY):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.20)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(4.5)
    r1 = p.add_run(title + " ")
    font_run(r1, lang, 9.5, bold=True, color=color)
    r2 = p.add_run(body)
    font_run(r2, lang, 9.5)
    return p


def source(doc, items, lang, prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(4)
    lead = prefix or ("Sources: " if lang == "en" else "资料来源：")
    r = p.add_run(lead)
    font_run(r, lang, 7.8, italic=True, color=GREY)
    for i, (label, url) in enumerate(items):
        if i:
            r = p.add_run("; ")
            font_run(r, lang, 7.8, italic=True, color=GREY)
        add_hyperlink(p, label, url, lang, 7.8)
    r = p.add_run(".")
    font_run(r, lang, 7.8, italic=True, color=GREY)
    return p


def figure(doc, idx, caption, lang, sources, width=6.15):
    pcap = para(doc, f"Figure {idx}. {caption}" if lang == "en" else f"图 {idx}：{caption}", lang, bold=True, size=8.7, color=NAVY, after=2)
    pcap.paragraph_format.keep_with_next = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    shape = p.add_run().add_picture(str(CHARTS[idx]), width=Inches(width))
    shape._inline.docPr.set("descr", caption)
    source(doc, sources, lang)


def table(doc, headers, rows, lang, widths, variance_col=None):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = tbl.rows[0].cells[i]
        cell.text = h
        shade(cell, NAVY)
        for r in cell.paragraphs[0].runs:
            font_run(r, lang, 8.2, bold=True, color=WHITE)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    repeat_header(tbl.rows[0])
    for row_i, vals in enumerate(rows):
        cells = tbl.add_row().cells
        for j, val in enumerate(vals):
            cells[j].text = str(val)
            if row_i % 2:
                shade(cells[j], LIGHT)
            for r in cells[j].paragraphs[0].runs:
                c = BLACK
                if variance_col == j:
                    c = GREEN if str(val).startswith("+") else RED if str(val).startswith("-") else BLACK
                font_run(r, lang, 8.2, bold=(j == 0), color=c)
            cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
    set_table_fixed(tbl, widths)
    return tbl


def new_page(doc):
    doc.add_page_break()


def title_page(doc, lang):
    if lang == "en":
        para(doc, "SK hynix Inc. (Nasdaq: SKHY)", lang, bold=True, size=18, color=NAVY, after=1, align=WD_ALIGN_PARAGRAPH.CENTER, heading_font=True)
        para(doc, "Q2 FY2026 EARNINGS UPDATE", lang, bold=True, size=14, color=BLUE, after=2, align=WD_ALIGN_PARAGRAPH.CENTER, heading_font=True)
        para(doc, "AI memory economics surge; headline earnings require normalization", lang, bold=True, size=11.5, color=ORANGE, after=8, align=WD_ALIGN_PARAGRAPH.CENTER, heading_font=True)
        meta = [
            "Published: 2 August 2026 | Delayed post-earnings review (T+4 calendar days)",
            "Initial post-IPO view: OUTPERFORM | Price target: US$175",
            "SKHY close (31 July 2026): US$143.73 | Implied upside: 21.8%",
            "Market capitalization: approximately US$1.02tn | Since-listing range: US$124.80–194.80",
        ]
        section = "EARNINGS SCORECARD"
        headers = ["Metric", "Reported", "Pre-results consensus", "Variance"]
        rows = [
            ["Revenue", "KRW79.32tn", "KRW84.10tn", "-KRW4.78tn / -5.7%"],
            ["Operating profit", "KRW60.54tn", "KRW64.10tn", "-KRW3.56tn / -5.5%"],
            ["Operating margin", "76.3%", "76.2% implied", "+0.1ppt"],
        ]
    else:
        para(doc, "SK hynix Inc.（纳斯达克：SKHY）", lang, bold=True, size=18, color=NAVY, after=1, align=WD_ALIGN_PARAGRAPH.CENTER, heading_font=True)
        para(doc, "2026 财年第二季度业绩更新报告", lang, bold=True, size=14, color=BLUE, after=2, align=WD_ALIGN_PARAGRAPH.CENTER, heading_font=True)
        para(doc, "AI 存储盈利能力跃升；报表利润仍须归一化", lang, bold=True, size=11.5, color=ORANGE, after=8, align=WD_ALIGN_PARAGRAPH.CENTER, heading_font=True)
        meta = [
            "发布日期：2026 年 8 月 2 日｜延迟财报点评（公告后第 4 个日历日）",
            "上市后首次观点：增持｜目标价：175 美元",
            "SKHY 收盘价（2026 年 7 月 31 日）：143.73 美元｜潜在上涨空间：21.8%",
            "市值：约 1.02 万亿美元｜上市以来区间：124.80–194.80 美元",
        ]
        section = "业绩计分卡"
        headers = ["指标", "实际值", "财报前一致预期", "差异"]
        rows = [
            ["营业收入", "79.32 万亿韩元", "84.10 万亿韩元", "-4.78 万亿 / -5.7%"],
            ["营业利润", "60.54 万亿韩元", "64.10 万亿韩元", "-3.56 万亿 / -5.5%"],
            ["营业利润率", "76.3%", "隐含 76.2%", "+0.1 个百分点"],
        ]
    for line in meta:
        para(doc, line, lang, size=9.2, bold=line.startswith("Initial") or line.startswith("上市后"), after=1.5)
    heading(doc, section, 2)
    table(doc, headers, rows, lang, [2300, 2100, 2500, 2460], variance_col=3)
    source(doc, [("SK hynix 2Q26 release" if lang == "en" else "SK hynix 2Q26 财报公告", URLS["release"]), ("Yonhap pre-results consensus" if lang == "en" else "韩联社财报前一致预期", URLS["consensus"])], lang)

    heading(doc, "INVESTMENT TAKEAWAYS" if lang == "en" else "投资要点", 2)
    if lang == "en":
        bullet(doc, "A modest top-line/operating-profit miss against an unusually elevated bar.", "Revenue and operating profit missed the averages of 14 local brokerages by 5.7% and 5.5%, respectively. The miss does not invalidate the cycle: revenue still grew 51% sequentially and 257% year over year, while operating margin expanded to 76%.", lang)
        bullet(doc, "The quality of operating earnings was high; the quality of reported net income was not.", "Gross margin reached 83%, but KRW93.92tn of net profit included KRW63.27tn of investment-asset gains. We remove these gains, on an illustrative after-tax basis, from our normalized earnings and valuation.", lang)
        bullet(doc, "HBM4, long-term agreements and leading-edge NAND support a longer, less volatile upcycle.", "HBM4 shipments began in Q2, HBM4E samples were delivered in the first half, and roughly ten customer LTAs now include pricing, volatility and deposit mechanisms. These improve visibility, although they cannot eliminate memory cyclicality.", lang)
        bullet(doc, "Initial post-IPO OUTPERFORM; US$175 target.", "Our target rounds a 7.5x multiple on FY2027E normalized ADR EPS of US$23.12. The discount to broader AI infrastructure leaders reflects capital intensity, supply-response risk, governance and ADR/FX complexity.", lang)
    else:
        bullet(doc, "收入和营业利润小幅不及极高预期门槛。", "营业收入和营业利润分别较 14 家韩国券商平均预期低 5.7% 和 5.5%。但周期逻辑未被破坏：收入环比增长 51%、同比增长 257%，营业利润率升至 76%。", lang)
        bullet(doc, "经营利润质量高，但报表净利润质量偏低。", "毛利率达到 83%，然而 93.92 万亿韩元净利润中包含 63.27 万亿韩元投资资产收益。我们的归一化盈利和估值按示意税后口径剔除该收益。", lang)
        bullet(doc, "HBM4、长期协议与先进 NAND 有望拉长并平滑上行周期。", "HBM4 已在第二季度出货，HBM4E 样品在上半年交付，约十份客户长期协议纳入价格、波动与定金机制。这些安排提高可见度，但无法消除存储行业周期性。", lang)
        bullet(doc, "上市后首次给予“增持”，目标价 175 美元。", "目标价对应 2027 财年归一化 ADR 每股收益 23.12 美元的 7.5 倍市盈率并取整。相对其他 AI 基础设施龙头的折价反映资本强度、供给反应、治理以及 ADR/汇率复杂性。", lang)


def results_pages(doc, lang):
    new_page(doc)
    heading(doc, "REVENUE: AI MEMORY PRICING POWER" if lang == "en" else "收入：AI 存储定价权释放", 1)
    if lang == "en":
        para(doc, "Q2 revenue reached KRW79.32tn, up 51% QoQ and 257% YoY. The step-up was driven by both volume and price, with DRAM bit shipments growing at a high-single-digit rate and ASP rising approximately 30% sequentially. NAND bit shipments increased in the mid-teens while ASP rose in the mid-50% range, producing an exceptional pricing contribution. The result fell short of the KRW84.1tn pre-release brokerage average, but it still represented the strongest quarter in the eight-quarter history shown below.", lang)
    else:
        para(doc, "第二季度收入达到 79.32 万亿韩元，环比增长 51%、同比增长 257%。量价共同推动收入跃升：DRAM 位元出货量环比高个位数增长，平均售价约上涨 30%；NAND 位元出货量环比中十位数增长，平均售价上涨中 50% 区间，价格贡献尤为突出。实际收入低于财报前 84.1 万亿韩元券商平均预期，但仍是下图八个季度中的历史最高水平。", lang)
    figure(doc, 1, "Quarterly revenue accelerated sharply into 2Q26" if lang == "en" else "季度收入在 2026 年第二季度显著加速", lang, [("Company quarterly presentations" if lang == "en" else "公司历季业绩演示材料", URLS["presentation"])])
    if lang == "en":
        para(doc, "The revenue miss is best interpreted as a mismatch between a very steep expectation curve and delivery timing rather than as demand destruction. AI server memory remains supply constrained, customer qualification schedules are lumpy, and product mix can shift across quarter boundaries. Still, the gap matters because the IPO and the sharp pre-earnings rally had capitalized a near-perfect outcome. In the next two quarters, investors should monitor whether HBM4 volume conversion and NAND pricing allow revenue to reaccelerate without a disproportionate inventory build.", lang)
    else:
        para(doc, "收入不及预期更适合解释为陡峭预期曲线与交付时点错位，而非需求破坏。AI 服务器存储仍受供给约束，客户认证节奏并不平滑，产品组合也可能跨季度迁移。但该差异仍然重要，因为 IPO 以及财报前的急涨已经计入接近完美的结果。未来两个季度应重点观察 HBM4 放量和 NAND 涨价能否推动收入再加速，同时避免库存不成比例上升。", lang)
    figure(doc, 2, "Operating profit expanded faster than revenue" if lang == "en" else "营业利润增速快于收入", lang, [("SK hynix 2Q26 presentation" if lang == "en" else "SK hynix 2Q26 业绩演示", URLS["presentation"]), ("Prior quarterly presentations" if lang == "en" else "公司过往季度演示", URLS["q1"])])

    new_page(doc)
    heading(doc, "MARGINS AND EARNINGS QUALITY" if lang == "en" else "利润率与盈利质量", 1)
    if lang == "en":
        para(doc, "Operating profit of KRW60.54tn increased 61% QoQ and 557% YoY, and operating margin expanded four points sequentially to 76.3%. Gross profit reached KRW65.99tn and gross margin 83.2%, while cost of goods sold was only KRW13.33tn. The gap between gross and operating margin remained narrow: SG&A was KRW5.45tn, evidence that revenue growth is scaling through a highly fixed-cost manufacturing platform. EBITDA was KRW64.56tn, equal to an 81% margin.", lang)
    else:
        para(doc, "营业利润为 60.54 万亿韩元，环比增长 61%、同比增长 557%，营业利润率环比提升约 4 个百分点至 76.3%。毛利润为 65.99 万亿韩元，毛利率 83.2%，营业成本仅 13.33 万亿韩元。毛利率与营业利润率之间的差距依然较小，销售及管理费用为 5.45 万亿韩元，显示收入增长正通过高固定成本制造平台形成强大经营杠杆。EBITDA 为 64.56 万亿韩元，利润率 81%。", lang)
    figure(doc, 3, "Gross and operating margins reached new highs" if lang == "en" else "毛利率与营业利润率创出新高", lang, [("Company quarterly presentations" if lang == "en" else "公司历季业绩演示材料", URLS["presentation"])])
    if lang == "en":
        para(doc, "Reported net profit of KRW93.92tn exceeded operating profit because pre-tax income included KRW60.89tn of other non-operating profit, principally KRW63.27tn of investment-asset gains. The quarter therefore cannot be valued on reported EPS of KRW131,478. We estimate normalized Q2 net income at roughly KRW45.5tn by removing the investment gain after applying a 23.5% illustrative tax rate. We use the same method to normalize the KRW9.94tn investment valuation gain disclosed in Q1. This is our analytical adjustment—not a company-reported non-GAAP measure—and actual tax realization may differ materially.", lang)
    else:
        para(doc, "报表净利润 93.92 万亿韩元高于营业利润，原因是税前利润中包括 60.89 万亿韩元其他非经营收益，主要为 63.27 万亿韩元投资资产收益。因此不能用本季报表每股收益 131,478 韩元直接估值。我们按 23.5% 的示意税率剔除该投资收益，估算第二季度归一化净利润约 45.5 万亿韩元；第一季度披露的 9.94 万亿韩元投资估值收益也采用相同方法处理。该口径是我们的分析调整，并非公司披露的非 GAAP 指标，实际税务实现可能显著不同。", lang)
    figure(doc, 4, "Reported net income was inflated by investment gains" if lang == "en" else "投资收益显著抬高报表净利润", lang, [("SK hynix 2Q26 presentation, pp. 6–7" if lang == "en" else "SK hynix 2Q26 演示第 6–7 页", URLS["presentation"]), ("SK hynix 1Q26 presentation" if lang == "en" else "SK hynix 1Q26 演示", URLS["q1"])])

    new_page(doc)
    heading(doc, "EXPECTATIONS AND PRODUCT MIX" if lang == "en" else "预期差与产品组合", 1)
    if lang == "en":
        para(doc, "The reported miss was concentrated in scale, not profitability. Against the KRW84.1tn revenue and KRW64.1tn operating-profit averages published by Yonhap before results, SK hynix missed by KRW4.78tn and KRW3.56tn. The implied consensus operating margin was 76.2%, nearly identical to the 76.3% delivered. This pattern suggests the Street largely understood unit economics but overestimated shipment timing and/or mix. It also explains why the result can be operationally excellent while still disappointing the stock market.", lang)
    else:
        para(doc, "本季不及预期主要体现在规模，而非盈利能力。相对于韩联社在财报前公布的 84.1 万亿韩元收入和 64.1 万亿韩元营业利润平均预期，公司分别低 4.78 万亿和 3.56 万亿韩元。一致预期隐含营业利润率为 76.2%，与实际 76.3% 几乎相同。这表明市场基本判断对单位经济性，却高估了出货节奏或产品组合，也解释了为何经营表现极强、股价反应却可能失望。", lang)
    figure(doc, 5, "The quarter missed elevated revenue and operating-profit expectations" if lang == "en" else "收入与营业利润低于高企预期", lang, [("Yonhap, 26 July 2026" if lang == "en" else "韩联社，2026 年 7 月 26 日", URLS["consensus"]), ("Company release" if lang == "en" else "公司财报公告", URLS["release"])])
    if lang == "en":
        para(doc, "Product mix shifted toward NAND: DRAM declined to 73% of revenue from 78% in Q1, while NAND rose to 27% from 21%. This was not negative mix dilution because NAND pricing rose much faster than DRAM pricing and 321-layer products increased their share. Management expects the 321-layer generation to become the largest NAND production share and to account for roughly half of domestic NAND capacity by year-end. The mix shift diversifies the earnings engine but also raises sensitivity to enterprise SSD demand and the speed of industry NAND supply response.", lang)
    else:
        para(doc, "产品组合向 NAND 倾斜：DRAM 收入占比从第一季度的 78% 降至 73%，NAND 从 21% 升至 27%。这并非负面组合稀释，因为 NAND 涨价幅度远高于 DRAM，321 层产品占比也在提高。管理层预计 321 层产品将成为 NAND 最大生产占比，并在年底前占韩国 NAND 产能约一半。组合变化使盈利引擎更分散，但也提高了对企业级 SSD 需求和行业 NAND 供给反应速度的敏感度。", lang)
    figure(doc, 6, "NAND gained share as pricing accelerated" if lang == "en" else "NAND 在价格加速中提升收入占比", lang, [("Company quarterly presentations" if lang == "en" else "公司历季业绩演示材料", URLS["presentation"])])


def balance_outlook_pages(doc, lang):
    new_page(doc)
    heading(doc, "BALANCE SHEET AND CASH FLOW" if lang == "en" else "资产负债表与现金流", 1)
    if lang == "en":
        para(doc, "Cash and short-term investments rose to KRW87.96tn at quarter-end, while interest-bearing debt was KRW18.59tn, leaving approximately KRW69.4tn of net cash. This is a radical improvement from the net-debt position two years earlier and provides strategic capacity for HBM, advanced packaging and leading-edge NAND investment. The US ADR offering added substantial liquidity: the final prospectus shows 177.9m ADSs sold at US$149, each representing 0.1 common share, and approximately US$26.2bn of net proceeds.", lang)
    else:
        para(doc, "期末现金及短期投资升至 87.96 万亿韩元，有息负债为 18.59 万亿韩元，净现金约 69.4 万亿韩元。相较两年前的净负债状态，这是结构性改善，为 HBM、先进封装与领先 NAND 投资提供战略能力。美国 ADR 发行进一步补充流动性：最终招股书显示，公司以每份 149 美元发行 1.779 亿份 ADS，每份代表 0.1 股普通股，净募资约 262 亿美元。", lang)
    figure(doc, 7, "Liquidity expanded faster than debt" if lang == "en" else "流动性增速明显快于负债", lang, [("SK hynix 2Q26 presentation" if lang == "en" else "SK hynix 2Q26 业绩演示", URLS["presentation"]), ("SEC final prospectus" if lang == "en" else "SEC 最终招股书", URLS["prospectus"])])
    if lang == "en":
        para(doc, "Operating cash flow was KRW65.71tn and purchases of property, plant and equipment were KRW10.67tn, implying illustrative free cash flow of KRW55.04tn before other investing items. This exceptional conversion benefited from pricing, but working capital absorbed some cash: accounts receivable reached KRW47.82tn and inventory KRW17.99tn. Receivables are the more important watch item because rapid revenue growth and customer deposits can complicate quarter-to-quarter interpretation. We prefer to judge cash conversion over multiple quarters and against capital commitments, not from a single peak-margin period.", lang)
    else:
        para(doc, "经营现金流为 65.71 万亿韩元，购置物业、厂房及设备支出 10.67 万亿韩元，在不考虑其他投资项目之前，示意自由现金流约为 55.04 万亿韩元。强劲转化受益于价格上涨，但营运资本仍占用部分现金：应收账款升至 47.82 万亿韩元，存货为 17.99 万亿韩元。应收账款尤其值得跟踪，因为收入高速增长与客户定金可能使单季解读复杂。我们更倾向以多个季度和资本承诺为尺度判断现金转化，而非用单个峰值利润率季度外推。", lang)
    figure(doc, 8, "Cash generation remained strong after capital spending" if lang == "en" else "资本开支后现金创造能力仍然强劲", lang, [("SK hynix 2Q26 presentation, pp. 9–10" if lang == "en" else "SK hynix 2Q26 演示第 9–10 页", URLS["presentation"])])

    new_page(doc)
    heading(doc, "OUTLOOK, GUIDANCE AND EXECUTION" if lang == "en" else "展望、指引与执行重点", 1)
    if lang == "en":
        para(doc, "SK hynix did not provide formal quarterly revenue, EPS or margin guidance. It did provide operating indicators: Q3 DRAM bit shipments are expected to rise approximately 10% QoQ and NAND bits by a low-single-digit percentage. For full-year 2026, management expects DRAM demand growth in the mid-20% range and NAND growth in the high teens. These indicators support continued expansion, but they leave ASP and mix—the two largest profit sensitivities—to investor judgment.", lang)
    else:
        para(doc, "SK hynix 未提供正式的单季收入、每股收益或利润率指引，但给出了运营指标：第三季度 DRAM 位元出货量预计环比增长约 10%，NAND 位元出货量预计低个位数增长。管理层预计 2026 年全年 DRAM 需求同比增长中 20% 区间，NAND 增长高十位数区间。这些指标支持继续扩张，但仍把平均售价与产品组合这两项最大利润敏感因素留给投资者判断。", lang)
    figure(doc, 9, "Management expects memory demand to remain robust" if lang == "en" else "管理层预计存储需求保持强劲", lang, [("SK hynix 2Q26 presentation" if lang == "en" else "SK hynix 2Q26 业绩演示", URLS["presentation"]), ("Earnings-call replay" if lang == "en" else "业绩会音频回放", URLS["call"])])
    if lang == "en":
        bullet(doc, "HBM roadmap:", "HBM4 shipments began in Q2 and should ramp through the second half. HBM4E samples were delivered in the first half, preserving SK hynix's technology cadence as accelerator platforms move to greater bandwidth and power efficiency.", lang)
        bullet(doc, "Contract architecture:", "Approximately ten customer long-term agreements now use pricing, volatility-sharing and deposit structures. These mechanisms improve demand visibility and capital planning, yet they may defer rather than remove risk if end demand slows.", lang)
        bullet(doc, "Capacity and CapEx:", "2026 CapEx is expected in the high-KRW40tn range. M15X and Yongin Fab 1 are being accelerated toward early 2027, while P&T7 and M17 will ramp in phases. The balance sheet can fund this plan, but returns depend on maintaining discipline as competitors add supply.", lang)
        bullet(doc, "Product execution:", "SOCAMM2 using 1c-nanometer DRAM has begun, while 321-layer NAND is set to become the largest production generation. Qualification, yield and advanced-packaging throughput are the critical operating proof points.", lang)
        para(doc, "The company did not publish a written transcript with the results; this report uses the official presentation and call replay. Statements attributed to management are paraphrased rather than quoted.", lang, size=8.5, color=GREY)
    else:
        bullet(doc, "HBM 路线图：", "HBM4 已在第二季度出货，并将在下半年持续爬坡；HBM4E 样品已于上半年交付。随着加速器平台追求更高带宽与能效，公司技术节奏得以延续。", lang)
        bullet(doc, "合同架构：", "约十份客户长期协议引入定价、波动共担与定金结构。这些机制提升需求可见度和资本规划能力，但若终端需求转弱，风险可能只是推迟而非消失。", lang)
        bullet(doc, "产能与资本开支：", "2026 年资本开支预计处于 40 万亿韩元高段。M15X 与龙仁第一晶圆厂目标提前至 2027 年初，P&T7 和 M17 将分阶段爬坡。资产负债表足以支持计划，但回报仍取决于竞争对手扩产时的行业纪律。", lang)
        bullet(doc, "产品执行：", "采用 1c 纳米 DRAM 的 SOCAMM2 已启动，321 层 NAND 将成为最大量产世代。客户认证、良率和先进封装吞吐量是关键运营验证点。", lang)
        para(doc, "公司未随财报发布书面电话会逐字稿；本报告使用官方演示材料和电话会音频回放。所有管理层表述均为转述，不作直接引语。", lang, size=8.5, color=GREY)
    source(doc, [("Official call replay" if lang == "en" else "官方电话会回放", URLS["call"]), ("Company IR" if lang == "en" else "公司投资者关系网页", URLS["ir"])], lang)


def estimates_valuation_pages(doc, lang):
    new_page(doc)
    heading(doc, "ESTIMATE REVISIONS" if lang == "en" else "盈利预测调整", 1)
    if lang == "en":
        para(doc, "Because this is our initial post-IPO report, the 'old' column below is an external pre-results reference case from Hanwha Securities dated 22 June 2026, not a prior estimate published by us. Our new estimates incorporate the Q2 miss, the Q3 bit-shipment indicators, HBM4 ramp, current NAND pricing and a deliberate normalization of investment gains. We forecast Q3 revenue/operating profit of KRW92tn/KRW72tn and Q4 of KRW101tn/KRW78tn.", lang)
        headers = ["KRW tn except EPS", "FY26E old*", "FY26E new", "Change", "FY27E old*", "FY27E new", "Change"]
        rows = [
            ["Revenue", "339.4", "324.9", "-4.3%", "506.8", "420.0", "-17.1%"],
            ["Operating profit", "266.7", "248.2", "-6.9%", "419.5", "315.0", "-24.9%"],
            ["Normalized net income", "217.3", "190.2", "-12.5%", "319.0", "242.0", "-24.1%"],
            ["Normalized ADR EPS (US$)", "20.75", "18.16", "-12.5%", "30.46", "23.12", "-24.1%"],
        ]
    else:
        para(doc, "由于这是我们在 ADR 上市后的首份报告，下表“原预测”采用韩华证券 2026 年 6 月 22 日的财报前外部参考情景，并非我们此前发布的预测。新预测纳入第二季度不及预期、第三季度位元出货指标、HBM4 爬坡、当前 NAND 价格，并对投资收益作审慎归一化。我们预计第三季度收入/营业利润为 92/72 万亿韩元，第四季度为 101/78 万亿韩元。", lang)
        headers = ["万亿韩元，EPS 除外", "FY26E 原*", "FY26E 新", "变化", "FY27E 原*", "FY27E 新", "变化"]
        rows = [
            ["营业收入", "339.4", "324.9", "-4.3%", "506.8", "420.0", "-17.1%"],
            ["营业利润", "266.7", "248.2", "-6.9%", "419.5", "315.0", "-24.9%"],
            ["归一化净利润", "217.3", "190.2", "-12.5%", "319.0", "242.0", "-24.1%"],
            ["归一化 ADR EPS（美元）", "20.75", "18.16", "-12.5%", "30.46", "23.12", "-24.1%"],
        ]
    table(doc, headers, rows, lang, [2200, 1300, 1300, 1200, 1300, 1300, 1200], variance_col=3)
    source(doc, [("Hanwha pre-results research (old*)" if lang == "en" else "韩华证券财报前研究（原预测*）", URLS["hanwha"]), ("Company Q2 materials" if lang == "en" else "公司第二季度材料", URLS["presentation"])], lang)
    figure(doc, 10, "Our post-results estimates are below the external pre-results reference case" if lang == "en" else "财报后预测低于外部财报前参考情景", lang, [("Hanwha Securities" if lang == "en" else "韩华证券", URLS["hanwha"]), ("Our estimates" if lang == "en" else "本报告预测", URLS["release"])])
    if lang == "en":
        para(doc, "The largest reduction is FY2027 operating profit because the pre-results case extrapolated current scarcity economics unusually far. Our FY2027 model still assumes 29% revenue growth and a 75% operating margin, both exceptional. Forecast risk is two-sided: tighter supply and stronger HBM pricing could produce upside, while faster capacity additions, customer inventory digestion or a softer AI infrastructure cycle could cause margins to mean-revert much sooner.", lang)
        para(doc, "Our quarterly bridge assumes Q3 revenue of KRW92tn on approximately 10% DRAM bit growth, low-single-digit NAND bit growth and continued ASP support, followed by KRW101tn in Q4 as HBM4 volumes and year-end enterprise demand broaden the mix. Operating profit of KRW72tn in Q3 and KRW78tn in Q4 implies margins near 78% and 77%, respectively. Those levels are deliberately below a mechanical extrapolation of Q2 gross-margin expansion: we allow for launch costs, product-transition inefficiency, higher depreciation and a less favorable incremental price/cost spread as the base becomes larger.", lang)
        para(doc, "For FY2027, the model treats HBM as a structural premium product but not as permanently immune to negotiation. Revenue grows to KRW420tn, while operating profit reaches KRW315tn. The resulting 75% operating margin embeds continued technology leadership and favorable utilization, yet also assumes some normalization from quarterly peak conditions. Normalized net income of KRW242tn is derived from operating earnings and recurring below-the-line items rather than from carrying forward Q2 investment gains. This distinction is central to the valuation: a one-off mark-to-market benefit can strengthen book value and liquidity, but it should not receive an earnings multiple as if it were recurring semiconductor profit.", lang)
        para(doc, "Sensitivity is unusually high because a one-point change in FY2027 operating margin changes operating profit by roughly KRW4.2tn before tax. Volume and ASP assumptions also interact: stronger pricing may lower customer order urgency or attract supply, while faster bit growth may require lower blended prices. We therefore view the forecast as a disciplined center point rather than a precise prediction. The most useful checkpoints are quarterly HBM qualification and shipment disclosures, the relationship between inventory and sales, CapEx commitments, and whether customer deposits continue to rise alongside long-term agreements.", lang)
    else:
        para(doc, "最大下调来自 2027 财年营业利润，因为财报前参考情景把当前稀缺性经济延伸得异常久。我们的 2027 财年模型仍假设收入增长 29%、营业利润率 75%，两者都处于极高水平。预测风险双向存在：供给更紧和 HBM 涨价可能带来上行，而更快扩产、客户库存消化或 AI 基础设施周期转弱则可能使利润率更早均值回归。", lang)
        para(doc, "季度衔接方面，我们预计第三季度收入 92 万亿韩元，基础是假设 DRAM 位元出货增长约 10%、NAND 低个位数增长且平均售价继续获得支撑；第四季度随着 HBM4 放量和年末企业需求扩大，收入升至 101 万亿韩元。第三、第四季度营业利润分别为 72 万亿和 78 万亿韩元，对应利润率约 78% 和 77%。该水平低于对第二季度毛利率扩张的机械外推，因为我们计入新品发布成本、产品迁移低效、更高折旧以及收入基数增大后的增量价差收窄。", lang)
        para(doc, "2027 财年模型把 HBM 视为结构性溢价产品，但并不假设其永久免于价格谈判。收入增至 420 万亿韩元，营业利润 315 万亿韩元，75% 营业利润率既包含持续技术领先和高产能利用，也包含相对单季峰值的适度回归。归一化净利润 242 万亿韩元来自经营利润和经常性线下项目，不延续第二季度投资收益。这一区分是估值核心：一次性按市值计量收益可以增强账面价值和流动性，却不应像持续性半导体利润一样获得市盈率。", lang)
        para(doc, "敏感度很高：2027 财年营业利润率每变化 1 个百分点，税前营业利润约变化 4.2 万亿韩元。销量与平均售价也相互作用，更强价格可能降低客户下单紧迫性或吸引供给，更快位元增长也可能要求更低综合价格。因此预测应被视为纪律严谨的中心情景，而非精确点估计。最有价值的验证指标包括 HBM 认证与出货、库存与销售关系、资本开支承诺，以及客户定金能否随长期协议继续增长。", lang)

    new_page(doc)
    heading(doc, "VALUATION AND PRICE TARGET" if lang == "en" else "估值与目标价", 1)
    if lang == "en":
        para(doc, "SKHY closed at US$143.73 on 31 July, below the US$149 IPO price after trading as high as US$194.80. The price path reflects both the scarcity premium assigned to a newly listed AI-memory leader and the rapid reset when Q2 results missed elevated expectations. Because the ADR has only a short trading history, conventional 52-week statistics are not meaningful; all range references are since listing.", lang)
    else:
        para(doc, "SKHY 于 7 月 31 日收于 143.73 美元，低于 149 美元 IPO 价格，上市后最高曾达 194.80 美元。价格路径同时反映市场对新上市 AI 存储龙头赋予的稀缺性溢价，以及第二季度未达高预期后的快速重估。由于 ADR 交易历史很短，常规 52 周统计意义有限；本报告所有区间均指上市以来。", lang)
    figure(doc, 11, "SKHY retraced after the Q2 release" if lang == "en" else "第二季度财报后 SKHY 回吐涨幅", lang, [("Yahoo Finance market data via yfinance" if lang == "en" else "通过 yfinance 获取的 Yahoo Finance 市场数据", "https://finance.yahoo.com/quote/SKHY/"), ("SEC final prospectus" if lang == "en" else "SEC 最终招股书", URLS["prospectus"])])
    if lang == "en":
        para(doc, "Our US$175 price target is the rounded output of 7.5x FY2027E normalized ADR EPS of US$23.12, implying a base value of approximately US$173. The multiple is intentionally conservative relative to asset-light AI beneficiaries because SK hynix must fund wafer capacity and packaging, operates in a historically cyclical industry, reports in Korean won and exposes ADR holders to governance and currency complexity. It is also above traditional memory-cycle trough multiples because HBM leadership, customer LTAs and net cash improve duration and resilience.", lang)
        headers = ["Scenario", "FY27E ADR EPS", "P/E", "Value", "Interpretation"]
        rows = [
            ["Bear", "US$18.50", "6.0x", "US$111", "Earlier supply response / margin normalization"],
            ["Base", "US$23.12", "7.5x", "US$173", "Rounded price target: US$175"],
            ["Bull", "US$27.00", "9.0x", "US$243", "HBM scarcity and LTAs extend peak economics"],
        ]
    else:
        para(doc, "175 美元目标价来自 2027 财年归一化 ADR 每股收益 23.12 美元的 7.5 倍市盈率，基础测算约 173 美元并取整。该倍数相对轻资产 AI 受益者更保守，因为 SK hynix 需要持续投入晶圆与封装产能，身处历史上强周期行业，以韩元报告，并使 ADR 持有人承担治理与汇率复杂性；但又高于传统存储周期底部倍数，因为 HBM 领先、客户长期协议和净现金改善了盈利久期与韧性。", lang)
        headers = ["情景", "FY27E ADR EPS", "市盈率", "价值", "含义"]
        rows = [
            ["悲观", "18.50 美元", "6.0 倍", "111 美元", "供给更早反应、利润率回归"],
            ["基准", "23.12 美元", "7.5 倍", "173 美元", "取整目标价 175 美元"],
            ["乐观", "27.00 美元", "9.0 倍", "243 美元", "HBM 稀缺与长期协议延长高盈利"],
        ]
    table(doc, headers, rows, lang, [1500, 1800, 1200, 1400, 3960])
    figure(doc, 12, "Scenario valuation spans US$111–243 per ADS" if lang == "en" else "情景估值区间为每份 ADS 111–243 美元", lang, [("Company prospectus share data" if lang == "en" else "公司招股书股本数据", URLS["prospectus"]), ("Our estimates and valuation" if lang == "en" else "本报告预测与估值", URLS["release"])])


def thesis_sources_pages(doc, lang):
    new_page(doc)
    heading(doc, "INVESTMENT THESIS, CATALYSTS AND RISKS" if lang == "en" else "投资逻辑、催化剂与风险", 1)
    if lang == "en":
        heading(doc, "Three-pillar thesis", 2)
        bullet(doc, "1. Structural AI memory intensity.", "The memory content and bandwidth required per accelerator and AI server continue to rise. HBM4 shipments and HBM4E sampling place SK hynix at the center of this shift, while SOCAMM2 creates an adjacent opportunity in low-power, high-capacity server memory.", lang)
        bullet(doc, "2. Visibility and technology leadership.", "Customer LTAs, deposits and volatility-sharing clauses make demand planning more durable than in prior commodity cycles. Leading DRAM nodes, 321-layer NAND and advanced packaging provide a route to retain price and cost leadership, provided execution remains on schedule.", lang)
        bullet(doc, "3. Balance-sheet transformation.", "KRW69.4tn of net cash and very strong operating cash flow lower financing risk just as the company accelerates M15X, Yongin Fab 1 and packaging capacity. The stronger balance sheet supports both growth investment and resilience in a future downturn.", lang)
        heading(doc, "Catalyst calendar", 2)
        bullet(doc, "August–October 2026:", "Evidence of HBM4 volume ramp, 1c DRAM and SOCAMM2 qualification, plus 321-layer NAND mix progression. Q3 results should test the approximately 10% QoQ DRAM bit-growth indicator and low-single-digit NAND shipment growth.", lang)
        bullet(doc, "Second half 2026:", "Additional HBM4/HBM4E customer qualification, LTA disclosures and customer deposits may improve confidence in FY2027 visibility. AI accelerator launch schedules are the principal external timing variable.", lang)
        bullet(doc, "Early 2027:", "M15X and Yongin Fab 1 milestones will frame supply growth, capital efficiency and the sustainability of current margins. Phased P&T7 and M17 ramps will determine whether packaging remains a bottleneck or becomes a competitive advantage.", lang)
        heading(doc, "Key risks", 2)
        bullet(doc, "Memory-cycle and supply response:", "Today's margins are extraordinary. Aggressive capacity additions by SK hynix or competitors could pressure ASPs and compress operating leverage more rapidly than modeled.", lang, RED)
        bullet(doc, "AI demand concentration and customer bargaining power:", "HBM demand depends on a small number of accelerator and cloud ecosystems. Design changes, delayed deployments or stronger customer negotiating leverage could reduce volume or price realization.", lang, RED)
        bullet(doc, "Execution and capital intensity:", "HBM4 yields, advanced-packaging throughput, node transitions and fab schedules are technically complex. CapEx in the high-KRW40tn range raises depreciation and the cost of execution errors.", lang, RED)
        bullet(doc, "Accounting, valuation and market-structure risk:", "Investment gains distort reported net income; ADR trading history is short; KRW/USD moves affect translated EPS; and foreign-private-issuer reporting differs from a US domestic 10-Q cadence.", lang, RED)
        heading(doc, "What would change our view", 2)
        para(doc, "We would become more constructive if HBM4 qualification broadens across customers without a material yield penalty, customer LTAs convert to deposits and binding volumes, and inventory grows more slowly than revenue while free cash flow remains robust after the announced capacity build. A clearer demonstration that advanced packaging, rather than wafer supply, is the limiting factor would also support a higher multiple because it would reduce the probability of rapid commodity-style oversupply. Conversely, evidence that customers are double ordering, renegotiating price formulas or delaying accelerator deployments would challenge the duration of peak economics even if near-term reported revenue remains strong.", lang)
        para(doc, "The most important downside trigger is not a single quarterly miss; it is a change in the price/cost trajectory. A sustained sequential ASP decline alongside higher depreciation would compress margins from both directions. We would also revisit the rating if CapEx commitments moved materially above the high-KRW40tn range without corresponding customer commitments, if receivables or inventory repeatedly outgrew sales, or if HBM4E qualification slipped into later platform cycles. The US$175 target therefore carries explicit execution gates. Meeting them could justify migration toward the bull scenario; missing several at once would move valuation closer to the bear case.", lang)
        para(doc, "Rating framework: OUTPERFORM means expected total return above the analyst's sector benchmark over 12 months. The rating and target are scenario-based opinions, not guarantees. Investors should size positions for memory-cycle volatility and newly listed ADR liquidity dynamics.", lang, size=8.5, color=GREY)
    else:
        heading(doc, "三大投资支柱", 2)
        bullet(doc, "1. AI 存储强度结构性上升。", "每颗加速器和每台 AI 服务器所需的存储容量与带宽持续提高。HBM4 出货和 HBM4E 送样使 SK hynix 位于这一趋势中心，SOCAMM2 则开辟低功耗、高容量服务器存储的相邻机会。", lang)
        bullet(doc, "2. 可见度与技术领先。", "客户长期协议、定金和波动共担条款使需求规划较以往商品周期更稳定。领先 DRAM 节点、321 层 NAND 和先进封装提供保持价格与成本优势的路径，前提是执行按计划推进。", lang)
        bullet(doc, "3. 资产负债表转型。", "69.4 万亿韩元净现金和强劲经营现金流降低融资风险，恰逢公司加速 M15X、龙仁第一晶圆厂与封装产能。更强资产负债表同时支持增长投资和未来下行周期韧性。", lang)
        heading(doc, "催化剂日历", 2)
        bullet(doc, "2026 年 8–10 月：", "HBM4 放量、1c DRAM 与 SOCAMM2 认证，以及 321 层 NAND 组合提升的证据。第三季度业绩将验证 DRAM 位元出货环比约增 10% 和 NAND 低个位数增长指标。", lang)
        bullet(doc, "2026 年下半年：", "更多 HBM4/HBM4E 客户认证、长期协议披露与客户定金可能提高 2027 财年可见度。AI 加速器发布时间表是主要外部时点变量。", lang)
        bullet(doc, "2027 年初：", "M15X 和龙仁第一晶圆厂里程碑将决定供给增长、资本效率和当前利润率的可持续性。P&T7 与 M17 分阶段爬坡将决定封装是瓶颈还是竞争优势。", lang)
        heading(doc, "主要风险", 2)
        bullet(doc, "存储周期与供给反应：", "当前利润率异常高。公司或竞争对手激进扩产可能压低平均售价，使经营杠杆比预测更快反转。", lang, RED)
        bullet(doc, "AI 需求集中与客户议价权：", "HBM 需求依赖少数加速器和云生态。设计变化、部署延迟或客户议价权增强可能降低销量或价格实现。", lang, RED)
        bullet(doc, "执行与资本强度：", "HBM4 良率、先进封装吞吐、节点迁移和晶圆厂进度技术复杂。40 万亿韩元高段资本开支增加折旧和执行失误成本。", lang, RED)
        bullet(doc, "会计、估值与市场结构：", "投资收益扭曲报表净利润；ADR 交易历史短；韩元兑美元影响换算 EPS；外国私人发行人披露节奏也不同于美国本土公司的 10-Q。", lang, RED)
        heading(doc, "哪些因素会改变观点", 2)
        para(doc, "若 HBM4 在更多客户中完成认证且良率代价有限，客户长期协议转化为定金和有约束力的采购量，库存增速低于收入，同时已宣布扩产后的自由现金流仍然强劲，我们会进一步上调观点。若能更清楚证明先进封装而非晶圆供给才是限制因素，也支持更高估值倍数，因为这降低商品式快速过剩概率。反之，若客户重复下单、重议价格公式或延迟加速器部署，即使近期报表收入仍强，也会削弱高盈利持续时间。", lang)
        para(doc, "最重要的下行触发因素并非单季不及预期，而是价格/成本轨迹发生变化。若平均售价持续环比下降，同时折旧上升，利润率将受到双向挤压。若资本开支在缺乏相应客户承诺时显著高于 40 万亿韩元高段，应收账款或库存连续快于收入增长，或 HBM4E 认证推迟到更晚平台周期，我们也会重新评估评级。因此 175 美元目标价带有明确执行门槛；达成这些门槛可使估值向乐观情景迁移，同时错过多项则会使估值接近悲观情景。", lang)
        para(doc, "评级定义：“增持”表示未来 12 个月预期总回报高于分析师所选行业基准。评级与目标价是基于情景的观点，并非保证。投资者应按存储周期波动和新上市 ADR 流动性特征控制仓位。", lang, size=8.5, color=GREY)

    new_page(doc)
    heading(doc, "SOURCES, METHODOLOGY AND DISCLOSURES" if lang == "en" else "资料来源、方法与披露", 1)
    if lang == "en":
        para(doc, "Primary-source hierarchy", lang, bold=True, size=10.5, color=NAVY)
        para(doc, "Financial statements, operating metrics and management outlook are drawn from SK hynix's 29 July 2026 release, official 2Q26 presentation and official earnings-call replay. Historical quarterly series use the corresponding company presentations. The SEC Form 6-K confirms the call schedule; the final 424B4 prospectus supplies ADS ratio, offering size, price, post-offering share count and net-proceeds data. No official written call transcript was available when this report was prepared.", lang)
        para(doc, "This report was prepared four calendar days after the release, outside the ideal same-day or 48-hour earnings-update window. The delay allowed incorporation of the first two full post-results trading sessions and the final prospectus share data, but market prices and consensus can change quickly. Readers should treat the stated price, market capitalization, foreign-exchange rate and implied upside as a 31 July 2026 snapshot. Operating and financial data remain anchored to the quarter ended 30 June 2026 unless a later date is explicitly identified.", lang)
        para(doc, "Consensus and market data", lang, bold=True, size=10.5, color=NAVY)
        para(doc, "Pre-results consensus is the average of 14 local brokerages reported by Yonhap on 26 July 2026. The external 'old estimate' reference is Hanwha Securities research dated 22 June 2026 and is not our prior coverage. SKHY price, market capitalization and trading-range data were fetched dynamically through yfinance after the 31 July close. The ADR's short history means the displayed range is since listing, not a full 52 weeks.", lang)
        para(doc, "Analytical conventions", lang, bold=True, size=10.5, color=NAVY)
        para(doc, "All financial figures are Korean won unless noted. Totals may not add because of rounding. Normalized net income removes disclosed Q1 and Q2 investment-asset valuation gains using an illustrative 23.5% tax rate; it is our estimate and not a company-reported non-GAAP measure. ADR EPS uses 728.8655m post-offering common shares, an ADS ratio of 0.1 common share per ADS and KRW1,436.6 per US dollar. Free cash flow is operating cash flow less purchases of property, plant and equipment and therefore excludes other investing flows. Price target upside excludes dividends.", lang)
        para(doc, "Important disclosure", lang, bold=True, size=10.5, color=NAVY)
        para(doc, "This report is independent analytical research prepared for informational purposes only. It is not investment advice, an offer, a solicitation or a fiduciary recommendation. Forecasts and valuation rely on assumptions that may prove incorrect. Memory semiconductors are cyclical, capital intensive and exposed to technology, geopolitical, customer-concentration, currency and supply-chain risks. Readers should review primary filings and obtain professional advice appropriate to their circumstances.", lang)
        refs = [("SK hynix 2Q26 earnings release", URLS["release"]), ("Official 2Q26 presentation", URLS["presentation"]), ("Official earnings-call replay", URLS["call"]), ("SK hynix investor relations", URLS["ir"]), ("SEC Form 6-K", URLS["sec6k"]), ("SEC final 424B4 prospectus", URLS["prospectus"]), ("Yonhap pre-results consensus", URLS["consensus"]), ("Hanwha Securities pre-results research", URLS["hanwha"]), ("SK hynix 1Q26 presentation", URLS["q1"]), ("SK hynix 4Q25 presentation", URLS["q425"]), ("SK hynix 3Q25 presentation", URLS["q325"]), ("SK hynix 2Q25 presentation", URLS["q225"]), ("SK hynix 1Q25 presentation", URLS["q125"]), ("SK hynix 4Q24 presentation", URLS["q424"]), ("SK hynix 3Q24 presentation", URLS["q324"])]
    else:
        para(doc, "一手资料层级", lang, bold=True, size=10.5, color=NAVY)
        para(doc, "财务报表、运营指标与管理层展望来自 SK hynix 2026 年 7 月 29 日财报公告、官方第二季度业绩演示和官方电话会音频回放。历史季度序列采用公司对应季度演示材料。SEC Form 6-K 用于确认电话会安排，最终 424B4 招股书提供 ADS 比例、发行规模、价格、发行后股本和净募资数据。报告准备时，公司未提供官方书面电话会逐字稿。", lang)
        para(doc, "本报告在财报发布后第 4 个日历日完成，晚于理想的当日或 48 小时财报更新窗口。延迟使报告能够纳入财报后两个完整交易日和最终招股书股本数据，但市场价格与一致预期会快速变化。所列股价、市值、汇率和潜在上涨空间应视为 2026 年 7 月 31 日快照；除非明确标注更晚日期，运营与财务数据均对应截至 2026 年 6 月 30 日的季度。", lang)
        para(doc, "一致预期与市场数据", lang, bold=True, size=10.5, color=NAVY)
        para(doc, "财报前一致预期采用韩联社 2026 年 7 月 26 日报道的 14 家韩国券商平均值。“原预测”外部参考来自韩华证券 2026 年 6 月 22 日研究，并非我们此前覆盖。SKHY 股价、市值和交易区间在 7 月 31 日收盘后通过 yfinance 动态获取。由于 ADR 历史很短，展示区间是上市以来而非完整 52 周。", lang)
        para(doc, "分析口径", lang, bold=True, size=10.5, color=NAVY)
        para(doc, "除特别说明外，所有财务数据均以韩元计，合计数可能因四舍五入存在差异。归一化净利润按 23.5% 示意税率剔除公司披露的第一、第二季度投资资产估值收益；这是本报告估算，并非公司非 GAAP 指标。ADR EPS 使用发行后 7.288655 亿股普通股、每份 ADS 代表 0.1 股普通股以及 1 美元兑 1,436.6 韩元。自由现金流定义为经营现金流减购置物业、厂房及设备，不含其他投资现金流。目标价上涨空间不含股息。", lang)
        para(doc, "重要披露", lang, bold=True, size=10.5, color=NAVY)
        para(doc, "本报告为独立分析研究，仅供信息参考，不构成投资建议、要约、招揽或受托建议。预测与估值依赖可能被证伪的假设。存储半导体具有周期性和资本密集特征，并面临技术、地缘政治、客户集中、汇率及供应链风险。读者应核阅一手披露，并根据自身情况获取专业意见。", lang)
        refs = [("SK hynix 2Q26 财报公告", URLS["release"]), ("官方 2Q26 业绩演示", URLS["presentation"]), ("官方电话会回放", URLS["call"]), ("SK hynix 投资者关系", URLS["ir"]), ("SEC Form 6-K", URLS["sec6k"]), ("SEC 最终 424B4 招股书", URLS["prospectus"]), ("韩联社财报前一致预期", URLS["consensus"]), ("韩华证券财报前研究", URLS["hanwha"]), ("SK hynix 1Q26 演示", URLS["q1"]), ("SK hynix 4Q25 演示", URLS["q425"]), ("SK hynix 3Q25 演示", URLS["q325"]), ("SK hynix 2Q25 演示", URLS["q225"]), ("SK hynix 1Q25 演示", URLS["q125"]), ("SK hynix 4Q24 演示", URLS["q424"]), ("SK hynix 3Q24 演示", URLS["q324"])]
    heading(doc, "REFERENCE LINKS" if lang == "en" else "参考链接", 2)
    for label, url in refs:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(2)
        add_hyperlink(p, label, url, lang, 8.5)


def build(lang):
    doc = setup_doc(lang)
    title_page(doc, lang)
    results_pages(doc, lang)
    balance_outlook_pages(doc, lang)
    estimates_valuation_pages(doc, lang)
    thesis_sources_pages(doc, lang)
    filename = "SKHY_Q2_FY2026_Earnings_Update.docx" if lang == "en" else "SKHY_Q2_FY2026_业绩更新报告_中文版.docx"
    path = OUT / filename
    doc.core_properties.title = "SK hynix Q2 FY2026 Earnings Update" if lang == "en" else "SK hynix 2026 财年第二季度业绩更新报告"
    doc.core_properties.subject = "Post-earnings equity research"
    doc.core_properties.author = "SKHY Research"
    doc.save(path)
    print(path)


if __name__ == "__main__":
    missing = [str(p) for p in CHARTS.values() if not p.exists()]
    if missing:
        raise FileNotFoundError("Missing charts: " + ", ".join(missing))
    build("en")
    build("cn")
